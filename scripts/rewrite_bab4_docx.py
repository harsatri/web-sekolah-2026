from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SOURCE_DOC = Path(r"C:\Users\upgra\Documents\trae_projects\WebSekolah-2026\working-bab4.docx")
OUTPUT_DOC = Path(r"C:\Users\upgra\Documents\trae_projects\WebSekolah-2026\working-bab4-revisi.docx")
FINAL_DOC = Path(r"C:\Users\upgra\OneDrive\Documents\2211103073_Harsa Tri Novenda_TA RAD BAB 4_revisi.docx")


def paragraph_text(element, doc):
    if element.tag.endswith("}p"):
        return Paragraph(element, doc).text.strip()
    return ""


def insert_paragraph_after(paragraph, text="", style=None, align=None, bold=False):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        run.bold = bold
    if align is not None:
        new_para.alignment = align
    return new_para


def add_paragraphs(anchor, texts, style="Normal"):
    current = anchor
    for text in texts:
        current = insert_paragraph_after(current, text, style=style)
    return current


def add_table_after(paragraph, data, style_name="Normal Table", font_size=10, left_align_cols=None):
    left_align_cols = set(left_align_cols or [])
    rows = len(data)
    cols = len(data[0]) if rows else 0
    table = paragraph._parent.add_table(rows=rows, cols=cols)
    table.style = style_name
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_idx, row in enumerate(data):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in left_align_cols else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.bold = True

    paragraph._p.addnext(table._tbl)
    return table


def add_table_section(anchor, caption, data, explanations, subtitle=None, font_size=10, left_align_cols=None):
    current = anchor
    if subtitle:
        current = insert_paragraph_after(current, subtitle, style="List Paragraph")
    current = insert_paragraph_after(current, caption, style="Caption")
    table = add_table_after(current, data, font_size=font_size, left_align_cols=left_align_cols)
    temp_p = OxmlElement("w:p")
    table._tbl.addnext(temp_p)
    current = Paragraph(temp_p, anchor._parent)
    if explanations:
        current.style = "Normal"
        current.add_run(explanations[0])
        for text in explanations[1:]:
            current = insert_paragraph_after(current, text, style="Normal")
    return current


def delete_between_markers(doc):
    body = doc.element.body
    children = list(body.iterchildren())
    start_idx = None
    end_idx = None

    for idx, child in enumerate(children):
        text = paragraph_text(child, doc)
        if text.startswith("BAB IV"):
            start_idx = idx
            continue
        if start_idx is not None and text == "DAFTAR PUSTAKA":
            end_idx = idx
            break

    if start_idx is None or end_idx is None:
        raise RuntimeError("Penanda BAB IV atau DAFTAR PUSTAKA tidak ditemukan.")

    for child in children[start_idx + 1:end_idx]:
        body.remove(child)

    return Paragraph(children[start_idx], doc)


def build_document():
    doc = Document(str(SOURCE_DOC))
    anchor = delete_between_markers(doc)

    anchor = insert_paragraph_after(anchor, "Pengumpulan dan Analisis Data", style="Heading 2")
    anchor = add_paragraphs(
        anchor,
        [
            "Pengumpulan data pada penelitian ini memadukan data primer dari wawancara lapangan dengan data sekunder dari dokumen sekolah, regulasi PPDB, statistik pendidikan, dan basis data aplikasi yang telah dibangun. Ruang lingkup pengambilan data difokuskan pada SD Negeri 1 Sokanegara dan SD Negeri 1 Kranji sebagai dua sekolah dasar negeri dengan tingkat peminat tinggi di wilayah Purwokerto Timur.",
            "Hasil wawancara menunjukkan bahwa proses seleksi awal calon siswa di sekolah objek penelitian masih berjalan secara semi-digital. Panitia telah memanfaatkan formulir daring dan rekap data, tetapi keputusan kelayakan tetap memerlukan verifikasi manual, diskusi internal, dan penafsiran kriteria yang belum sepenuhnya terstandar dalam satu sistem terintegrasi.",
            "Data aplikasi memperlihatkan bahwa sistem yang dibangun telah menyimpan empat submission riil pada tabel eligibility_submissions sampai tanggal 5 Mei 2026. Data tersebut menjadi sumber utama pembahasan pada Bab 4 karena dapat menunjukkan hubungan langsung antara kebutuhan lapangan, implementasi logika SAW, dan hasil rekomendasi sekolah yang dihasilkan sistem.",
            "Kriteria yang digunakan dalam perhitungan SAW terdiri atas usia, domisili, nilai rapor TK, prestasi akademik, prestasi non-akademik, kelengkapan dokumen, dan kondisi ekonomi keluarga. Keterlibatan orang tua tetap dikumpulkan sebagai data pendamping, tetapi variabel tersebut tidak dimasukkan ke dalam penjumlahan terbobot agar proses penilaian tetap fokus pada kriteria yang dapat diskor secara konsisten.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Analisis Sistem", style="Heading 3")
    anchor = insert_paragraph_after(anchor, "Analisis Sistem Existing", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Proses existing di sekolah objek penelitian dimulai dari pengisian data calon siswa oleh orang tua atau wali melalui formulir yang disediakan panitia. Data yang masuk kemudian direkap kembali oleh tim sekolah untuk diverifikasi satu per satu sebelum dibahas pada tahap penentuan kelayakan.",
            "Model kerja tersebut sudah membantu pengumpulan data dasar, tetapi belum mampu menghasilkan umpan balik kelayakan secara langsung. Panitia masih perlu menafsirkan usia, domisili, rapor, prestasi, dan kelengkapan dokumen secara terpisah sehingga peluang terjadinya perbedaan keputusan antarpetugas masih tetap ada.",
            "Kendala utama yang muncul pada alur existing adalah duplikasi pengisian, proses verifikasi yang memerlukan waktu cukup panjang, dan belum tersedianya dasar perhitungan kuantitatif yang dapat dijelaskan kembali kepada pihak luar. Kondisi tersebut berdampak pada rendahnya transparansi karena calon pendaftar belum memperoleh alasan kelayakan dalam bentuk skor yang mudah dipahami.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Analisis Sistem yang Diusulkan", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Sistem yang diusulkan pada penelitian ini berbentuk aplikasi website yang menyediakan proses pre-screening kelayakan calon siswa secara mandiri. Pengguna umum mengisi biodata, domisili, rapor TK, prestasi, dokumen, dan kondisi ekonomi, lalu sistem mengubah seluruh data tersebut ke dalam skor terstandar agar dapat dihitung dengan metode SAW.",
            "Alur sistem usulan dirancang lebih ringkas dibandingkan proses existing. Form multi-langkah membantu pengguna mengisi data secara bertahap, perhitungan SAW dijalankan secara otomatis saat form selesai, dan hasil akhir ditampilkan dalam bentuk nilai preferensi serta rekomendasi sekolah yang paling sesuai dengan kondisi calon siswa.",
            "Pendekatan ini memberi manfaat langsung bagi dua pihak. Panitia sekolah memperoleh data hasil seleksi awal yang sudah tersusun dalam format terukur, sedangkan orang tua atau pendamping dari TK memperoleh gambaran peluang kelayakan anak sebelum proses PPDB resmi dijalankan oleh sekolah dasar tujuan.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Perhitungan Metode Simple Additive Weighting (SAW)", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Perhitungan SAW pada penelitian ini dilakukan setelah seluruh data calon siswa dikonversi ke dalam skor numerik yang seragam. Struktur penilaian tersebut dibangun agar setiap kriteria memiliki pengaruh yang jelas terhadap hasil akhir, sekaligus memudahkan panitia menjelaskan alasan mengapa seorang calon siswa lebih direkomendasikan pada sekolah tertentu.",
            "Data uji yang digunakan pada pembahasan ini berasal dari submission riil aplikasi pada 4 Mei 2026 pukul 04.52 WIB yang telah dianonimkan menjadi Calon Siswa A. Alternatif yang dinilai adalah dua sekolah objek penelitian, yaitu SD Negeri 1 Sokanegara dan SD Negeri 1 Kranji, sehingga proses perhitungan dapat langsung menggambarkan cara kerja sistem saat dipakai pada lingkungan nyata.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Kriteria dan Bobot", style="Heading 4")

    table_41 = [
        ["No", "Kode", "Nama Kriteria", "Atribut", "Bobot", "Persentase"],
        ["1", "C1", "Usia", "Benefit", "0,30", "30%"],
        ["2", "C2", "Domisili", "Benefit", "0,25", "25%"],
        ["3", "C3", "Nilai rapor TK", "Benefit", "0,15", "15%"],
        ["4", "C4", "Prestasi akademik", "Benefit", "0,10", "10%"],
        ["5", "C5", "Prestasi non-akademik", "Benefit", "0,10", "10%"],
        ["6", "C6", "Kelengkapan dokumen", "Benefit", "0,07", "7%"],
        ["7", "C7", "Kondisi ekonomi", "Cost", "0,03", "3%"],
    ]
    expl_41 = [
        "Tabel 4.1 menunjukkan bahwa usia dan domisili ditempatkan sebagai kriteria dominan dalam sistem. Komposisi tersebut konsisten dengan hasil wawancara lapangan yang menegaskan bahwa sekolah dasar negeri tetap menempatkan kesiapan usia dan kedekatan wilayah sebagai pertimbangan utama pada tahap seleksi awal.",
        "Bobot pada tabel ini ditetapkan dari kebutuhan sekolah dan hasil diskusi dengan stakeholder, lalu digunakan oleh metode SAW pada tahap penjumlahan terbobot. Posisi SAW dalam penelitian ini bukan untuk membentuk bobot dari nol, melainkan untuk mengolah bobot yang telah disepakati menjadi nilai preferensi yang objektif dan mudah dijelaskan.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.1 Kriteria dan Bobot Penilaian",
        table_41,
        expl_41,
        font_size=9,
        left_align_cols={2},
    )

    table_42 = [
        ["No", "Rentang Usia per 1 Juli 2026", "Skor", "Keterangan"],
        ["1", "6,00 sampai 7,00 tahun", "5", "Rentang prioritas utama masuk SD."],
        ["2", "5,50 sampai kurang dari 6,00 tahun", "4", "Layak dengan verifikasi tambahan."],
        ["3", "Lebih dari 7,00 sampai 8,00 tahun", "3", "Masih dapat dipertimbangkan."],
        ["4", "Kurang dari 5,50 atau lebih dari 8,00 tahun", "0", "Tidak direkomendasikan oleh sistem."],
    ]
    expl_42 = [
        "Tabel 4.2 memperlihatkan bahwa skor tertinggi diberikan pada rentang usia 6 sampai 7 tahun. Rentang tersebut dipilih karena paling merepresentasikan usia masuk SD yang aman secara perkembangan dan paling sering dijadikan dasar prioritas oleh sekolah dasar negeri.",
        "Skor usia di bawah 6 tahun tetap disediakan agar sistem dapat menampung kasus khusus, tetapi nilainya dibuat lebih rendah untuk menandai perlunya verifikasi lanjutan. Skor nol diberikan pada usia yang berada di luar batas sistem agar hasil rekomendasi tidak menyesatkan pengguna sejak tahap awal pengisian.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.2 Subkriteria Usia",
        table_42,
        expl_42,
        subtitle="Subkriteria Usia",
        font_size=9,
        left_align_cols={1, 3},
    )

    table_43 = [
        ["No", "Kondisi Domisili", "Skor", "Dasar Penilaian"],
        ["1", "Satu kelurahan dengan sekolah", "5", "Prioritas tertinggi pada lingkungan terdekat."],
        ["2", "Satu kecamatan dengan sekolah", "3", "Masih berada pada wilayah layanan yang sama."],
        ["3", "Luar kecamatan", "1", "Prioritas paling rendah."],
    ]
    expl_43 = [
        "Tabel 4.3 menunjukkan bahwa domisili diperlakukan sebagai kriteria pembeda utama antaralternatif sekolah. Nilai domisili menjadi sangat penting karena dua sekolah objek penelitian berada pada kecamatan yang sama, tetapi tidak selalu berada pada kelurahan yang sama dengan alamat calon siswa.",
        "Skor ini menjadi penentu paling terlihat pada studi kasus Bab 4 karena seluruh kriteria lain Calon Siswa A bernilai sama untuk kedua sekolah. Perbedaan hasil akhir akhirnya ditentukan oleh fakta bahwa calon siswa berasal dari Kelurahan Sokanegara sehingga alternatif SD Negeri 1 Sokanegara memperoleh nilai domisili lebih tinggi.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.3 Subkriteria Domisili",
        table_43,
        expl_43,
        subtitle="Subkriteria Domisili",
        font_size=9,
        left_align_cols={1, 3},
    )

    table_44 = [
        ["No", "Aspek Perkembangan", "Kode", "Keterangan"],
        ["1", "Nilai agama dan moral", "AM", "Menilai akhlak dan kebiasaan positif anak."],
        ["2", "Sosial emosional", "SE", "Menilai kemampuan berinteraksi dan mengelola emosi."],
        ["3", "Fisik motorik", "FM", "Menilai koordinasi tubuh dan aktivitas motorik."],
        ["4", "Kognitif", "KOG", "Menilai kemampuan berpikir awal anak."],
        ["5", "Bahasa", "BHS", "Menilai komunikasi dan pemahaman instruksi."],
        ["6", "Seni", "SNI", "Menilai kreativitas dan ekspresi."],
        ["7", "Perilaku dan kemandirian", "PK", "Menilai kesiapan anak beraktivitas secara mandiri."],
    ]
    expl_44 = [
        "Tabel 4.4 memuat tujuh aspek rapor TK yang digunakan sistem untuk membentuk nilai C3. Pemilihan aspek dilakukan agar penilaian rapor tidak berhenti pada satu angka global, tetapi benar-benar mewakili perkembangan anak dari sisi akademik awal, sosial, dan kemandirian.",
        "Struktur aspek tersebut membuat proses input lebih konsisten bagi pengguna. Guru TK, orang tua, atau petugas yang membantu pengisian hanya perlu memilih kategori perkembangan pada setiap aspek, kemudian sistem menghitung rata-rata skor secara otomatis tanpa perhitungan manual.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.4 Aspek Penilaian Rapor TK",
        table_44,
        expl_44,
        subtitle="Subkriteria Nilai Rapor TK",
        font_size=9,
        left_align_cols={1, 3},
    )

    table_45 = [
        ["No", "Kategori Penilaian", "Rentang Rata-rata", "Skor", "Keterangan"],
        ["1", "BSB", "3,26 sampai 4,00", "4", "Berkembang sangat baik."],
        ["2", "BSH", "2,51 sampai 3,25", "3", "Berkembang sesuai harapan."],
        ["3", "MB", "1,76 sampai 2,50", "2", "Mulai berkembang."],
        ["4", "BB", "1,00 sampai 1,75", "1", "Belum berkembang."],
    ]
    expl_45 = [
        "Tabel 4.5 menunjukkan proses konversi kategori rapor TK ke dalam skor numerik 1 sampai 4. Konversi ini membuat data rapor yang semula bersifat kualitatif dapat langsung diolah sebagai bagian dari perhitungan SAW tanpa menghilangkan makna perkembangan anak.",
        "Nilai C3 pada sistem diperoleh dengan merata-ratakan seluruh skor aspek yang diisi. Mekanisme rata-rata ini membuat kekuatan dan kelemahan perkembangan anak tetap terlihat secara proporsional, bukan hanya ditentukan oleh satu aspek yang paling tinggi atau paling rendah.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.5 Konversi Skala Penilaian Rapor TK",
        table_45,
        expl_45,
        subtitle="Konversi Skala Penilaian Rapor TK",
        font_size=9,
        left_align_cols={1, 2, 4},
    )

    table_46 = [
        ["No", "Tingkat Prestasi", "Skor Akademik", "Skor Non-Akademik", "Contoh"],
        ["1", "Provinsi", "5", "5", "Juara lomba tingkat provinsi."],
        ["2", "Kabupaten/Kota", "4", "4", "Juara lomba tingkat kabupaten/kota."],
        ["3", "Kecamatan", "3", "3", "Juara lomba tingkat kecamatan."],
        ["4", "Tidak ada prestasi", "1", "1", "Tidak memiliki sertifikat prestasi."],
    ]
    expl_46 = [
        "Tabel 4.6 memperlihatkan bahwa prestasi akademik dan non-akademik diperlakukan sebagai dua kriteria yang berbeda, tetapi menggunakan struktur skor yang sama. Pendekatan ini dipilih agar sistem dapat membedakan capaian belajar formal dari capaian bakat atau kegiatan pengembangan diri.",
        "Nilai yang digunakan sistem adalah tingkat prestasi tertinggi yang dimiliki calon siswa pada masing-masing kategori. Cara tersebut membuat hasil perhitungan lebih sederhana, tetapi tetap cukup representatif untuk kebutuhan pre-screening pada tahap awal penerimaan siswa baru.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.6 Subkriteria Prestasi Akademik dan Non-Akademik",
        table_46,
        expl_46,
        subtitle="Subkriteria Prestasi Akademik dan Non-Akademik",
        font_size=9,
        left_align_cols={1, 4},
    )

    table_47 = [
        ["No", "Dokumen", "Status", "Keterangan"],
        ["1", "Kartu Keluarga", "Wajib", "Verifikasi identitas keluarga dan domisili."],
        ["2", "Akta kelahiran", "Wajib", "Verifikasi identitas dan usia anak."],
        ["3", "Rapor TK/PAUD", "Wajib", "Sumber data perkembangan anak."],
        ["4", "Foto 3x4", "Wajib", "Arsip dasar calon siswa."],
        ["5", "Surat rekomendasi psikolog", "Kondisional", "Wajib untuk usia 5,50 sampai kurang dari 6,00 tahun."],
    ]
    expl_47 = [
        "Tabel 4.7 menunjukkan bahwa kelengkapan dokumen tidak hanya dipakai sebagai formalitas administrasi, tetapi juga sebagai indikator kesiapan data calon siswa untuk diverifikasi sekolah. Keberadaan dokumen inti membantu panitia mengecek kecocokan identitas, usia, alamat, dan informasi perkembangan anak secara lebih cepat.",
        "Dokumen kondisional berupa surat rekomendasi psikolog disediakan untuk mengakomodasi kasus calon siswa yang usianya belum masuk rentang prioritas utama. Posisi dokumen ini penting karena sekolah tetap memerlukan dasar profesional ketika menerima anak yang usianya masih berada di bawah 6 tahun.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.7 Subkriteria Kelengkapan Dokumen",
        table_47,
        expl_47,
        subtitle="Subkriteria Kelengkapan Dokumen",
        font_size=9,
        left_align_cols={1, 3},
    )

    table_48 = [
        ["No", "Kondisi Kelengkapan", "Rumus", "Skor", "Keterangan"],
        ["1", "Seluruh dokumen wajib lengkap", "(4/4) x 5", "5,00", "Siap diverifikasi penuh."],
        ["2", "Tiga dokumen wajib lengkap", "(3/4) x 5", "3,75", "Masih memerlukan kelengkapan tambahan."],
        ["3", "Dua dokumen wajib lengkap", "(2/4) x 5", "2,50", "Administrasi belum stabil."],
        ["4", "Satu dokumen wajib lengkap", "(1/4) x 5", "1,25", "Administrasi sangat kurang."],
        ["5", "Tidak ada dokumen wajib lengkap", "(0/4) x 5", "0,00", "Belum layak diproses."],
    ]
    expl_48 = [
        "Tabel 4.8 menjelaskan bahwa skor C6 dibentuk dari proporsi jumlah dokumen yang tersedia dibandingkan jumlah dokumen yang wajib dipenuhi. Rumus ini membuat kelengkapan dokumen dapat dinilai secara objektif dan tidak berhenti pada keputusan biner lengkap atau tidak lengkap.",
        "Penyebut pada rumus dapat berubah menjadi lima apabila sistem mendeteksi bahwa surat rekomendasi psikolog menjadi dokumen wajib tambahan. Perubahan tersebut menjaga konsistensi karena setiap calon siswa dinilai berdasarkan paket dokumen yang memang relevan dengan kondisi usianya.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.8 Konversi Kelengkapan Dokumen",
        table_48,
        expl_48,
        subtitle="Konversi Kelengkapan Dokumen",
        font_size=9,
        left_align_cols={1, 2, 4},
    )

    table_49 = [
        ["No", "Rentang Penghasilan per Bulan", "Skor (xij)", "Normalisasi r7", "Keterangan"],
        ["1", "Rp0 sampai Rp500.000", "5", "0,20", "Kondisi ekonomi sangat terbatas."],
        ["2", "Rp500.001 sampai Rp1.000.000", "4", "0,25", "Kondisi ekonomi terbatas."],
        ["3", "Rp1.000.001 sampai Rp2.000.000", "3", "0,33", "Kondisi ekonomi menengah bawah."],
        ["4", "Rp2.000.001 sampai Rp3.000.000", "2", "0,50", "Kondisi ekonomi menengah."],
        ["5", "Lebih dari Rp3.000.000", "1", "1,00", "Kondisi ekonomi menengah atas."],
    ]
    expl_49 = [
        "Tabel 4.9 menunjukkan pengelompokan kondisi ekonomi ke dalam lima rentang penghasilan bulanan. Kriteria ini diberi bobot paling kecil karena fungsinya hanya sebagai faktor penyeimbang, bukan sebagai penentu utama seperti usia dan domisili.",
        "Posisi C7 sebagai kriteria cost membuat nilainya dinormalisasi secara terbalik pada tahap berikutnya. Pendekatan tersebut menjaga agar faktor ekonomi tetap masuk ke dalam keputusan sistem, tetapi tidak menggeser dominasi kriteria utama yang memang lebih erat dengan aturan penerimaan sekolah dasar negeri.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.9 Subkriteria Kondisi Ekonomi",
        table_49,
        expl_49,
        subtitle="Subkriteria Kondisi Ekonomi",
        font_size=9,
        left_align_cols={1, 4},
    )

    anchor = insert_paragraph_after(anchor, "Data Alternatif", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Data alternatif pada penelitian ini terdiri atas dua sekolah objek penelitian, yaitu SD Negeri 1 Sokanegara dan SD Negeri 1 Kranji. Nilai alternatif dibentuk dari submission riil Calon Siswa A pada 4 Mei 2026 yang disimpan di basis data aplikasi, kemudian nama siswa disamarkan agar identitas pribadi tetap terlindungi.",
            "Nilai C1, C3, C4, C5, C6, dan C7 pada kedua alternatif bernilai sama karena berasal dari profil calon siswa yang sama. Perbedaan nilai hanya muncul pada kriteria C2 karena alamat calon siswa berada di Kelurahan Sokanegara, sehingga kedekatan wilayah terhadap masing-masing sekolah tidak identik.",
        ],
    )

    table_410 = [
        ["Alternatif", "Sekolah", "C1", "C2", "C3", "C4", "C5", "C6", "C7"],
        ["A1", "SD Negeri 1 Sokanegara", "5", "5", "4,00", "3", "1", "5,00", "3"],
        ["A2", "SD Negeri 1 Kranji", "5", "3", "4,00", "3", "1", "5,00", "3"],
    ]
    expl_410 = [
        "Tabel 4.10 memperlihatkan matriks keputusan awal yang menjadi dasar pengolahan SAW. Nilai C3 sebesar 4,00 diperoleh dari rata-rata tujuh aspek rapor TK yang seluruhnya berada pada kategori sangat baik, sedangkan nilai C4 sebesar 3 berasal dari prestasi akademik tingkat kecamatan yang dimiliki calon siswa.",
        "Nilai C2 menjadi pembeda paling jelas pada tabel ini. Alternatif A1 memperoleh skor 5 karena sekolah berada pada kelurahan yang sama dengan domisili calon siswa, sedangkan alternatif A2 memperoleh skor 3 karena masih berada dalam kecamatan yang sama tetapi berbeda kelurahan.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.10 Data Alternatif Calon Siswa A",
        table_410,
        expl_410,
        font_size=9,
        left_align_cols={1},
    )

    anchor = insert_paragraph_after(anchor, "Normalisasi Matriks", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Normalisasi pada sistem dilakukan terhadap skala maksimum setiap kriteria yang telah dibakukan sejak awal. Kriteria C1, C2, C4, C5, dan C6 menggunakan pembagi 5, sedangkan C3 menggunakan pembagi 4 karena berasal dari rata-rata kategori rapor TK.",
            "Kriteria C7 diperlakukan sebagai cost sehingga normalisasinya dihitung dengan rumus 1/xij. Nilai normalisasi Calon Siswa A pada kedua alternatif adalah C1 = 5/5 = 1,00; C3 = 4,00/4 = 1,00; C4 = 3/5 = 0,60; C5 = 1/5 = 0,20; C6 = 5,00/5 = 1,00; dan C7 = 1/3 = 0,33.",
            "Perbedaan normalisasi hanya muncul pada kriteria domisili. Nilai r2 untuk alternatif A1 dihitung dari 5/5 = 1,00, sedangkan nilai r2 untuk alternatif A2 dihitung dari 3/5 = 0,60. Hasil tersebut memperlihatkan bahwa kedekatan alamat calon siswa menjadi faktor pembeda utama dalam perhitungan akhir.",
        ],
    )

    table_411 = [
        ["Alternatif", "r1", "r2", "r3", "r4", "r5", "r6", "r7"],
        ["A1", "1,00", "1,00", "1,00", "0,60", "0,20", "1,00", "0,33"],
        ["A2", "1,00", "0,60", "1,00", "0,60", "0,20", "1,00", "0,33"],
    ]
    expl_411 = [
        "Tabel 4.11 menunjukkan bahwa hampir seluruh nilai normalisasi kedua alternatif identik karena menggunakan profil calon siswa yang sama. Nilai r4 sebesar 0,60 berasal dari pembagian 3 terhadap 5, sedangkan nilai r5 sebesar 0,20 berasal dari pembagian 1 terhadap 5.",
        "Nilai r2 menjadi satu-satunya elemen yang membedakan kedua alternatif. Keberadaan selisih pada satu kriteria ini cukup untuk menggeser ranking akhir karena bobot domisili mencapai 25 persen dari total nilai preferensi.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.11 Hasil Normalisasi Matriks",
        table_411,
        expl_411,
        font_size=9,
    )

    anchor = insert_paragraph_after(anchor, "Perankingan", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Nilai preferensi akhir diperoleh dengan menjumlahkan hasil perkalian bobot dan nilai normalisasi pada setiap kriteria. Perhitungan alternatif A1 dilakukan dengan rumus V1 = (0,30 x 1,00) + (0,25 x 1,00) + (0,15 x 1,00) + (0,10 x 0,60) + (0,10 x 0,20) + (0,07 x 1,00) + (0,03 x 0,33) = 0,86.",
            "Perhitungan alternatif A2 dilakukan dengan rumus V2 = (0,30 x 1,00) + (0,25 x 0,60) + (0,15 x 1,00) + (0,10 x 0,60) + (0,10 x 0,20) + (0,07 x 1,00) + (0,03 x 0,33) = 0,76. Nilai V1 yang lebih besar menunjukkan bahwa SD Negeri 1 Sokanegara menjadi alternatif yang lebih sesuai untuk Calon Siswa A.",
        ],
    )

    table_412 = [
        ["Alternatif", "C1", "C2", "C3", "C4", "C5", "C6", "C7", "Nilai V", "Peringkat"],
        ["A1", "0,30", "0,25", "0,15", "0,06", "0,02", "0,07", "0,01", "0,86", "1"],
        ["A2", "0,30", "0,15", "0,15", "0,06", "0,02", "0,07", "0,01", "0,76", "2"],
    ]
    expl_412 = [
        "Tabel 4.12 memperlihatkan hasil penjumlahan kontribusi bobot untuk setiap alternatif. Alternatif A1 memperoleh nilai preferensi 0,86, sedangkan alternatif A2 memperoleh nilai 0,76. Selisih 0,10 tersebut sepenuhnya dipengaruhi oleh perbedaan kontribusi C2 pada kedua sekolah.",
        "Hasil perankingan menunjukkan bahwa SD Negeri 1 Sokanegara menjadi alternatif terbaik bagi Calon Siswa A. Nilai 0,86 pada sistem diklasifikasikan sebagai sangat direkomendasikan, sedangkan nilai 0,76 pada alternatif kedua masih berada pada kategori direkomendasikan tetapi tidak menjadi pilihan utama.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.12 Nilai Preferensi dan Perankingan",
        table_412,
        expl_412,
        font_size=9,
    )

    anchor = insert_paragraph_after(anchor, "Hasil Akhir dan Rekomendasi", style="Heading 4")
    anchor = add_paragraphs(
        anchor,
        [
            "Hasil akhir perhitungan menunjukkan bahwa sistem merekomendasikan Calon Siswa A untuk SD Negeri 1 Sokanegara. Keputusan tersebut konsisten dengan kondisi lapangan karena domisili calon siswa berada pada kelurahan yang sama dengan lokasi sekolah, sementara seluruh kriteria non-domisili tetap memenuhi standar kelayakan yang baik.",
            "Pola hasil tersebut menunjukkan bahwa metode SAW pada penelitian ini tidak hanya menghasilkan ranking, tetapi juga menyimpan jejak alasan keputusan secara transparan. Sekolah dapat menjelaskan hasil rekomendasi melalui nilai setiap kriteria, sedangkan pengguna memperoleh pemahaman yang lebih jelas mengenai posisi kelayakan anak sebelum proses PPDB resmi berlangsung.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Perancangan Sistem", style="Heading 2")
    anchor = insert_paragraph_after(anchor, "Tahap Requirements Planning", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Tahap requirements planning pada metode RAD dimulai dari pengumpulan kebutuhan pengguna melalui wawancara dengan panitia sekolah dan telaah terhadap alur PPDB yang sudah berjalan. Hasil tahap ini menunjukkan bahwa sistem harus mampu menampung dua kelompok pengguna utama, yaitu pengguna umum yang melakukan pre-screening dan admin sekolah yang memantau hasil rekomendasi.",
            "Kebutuhan fungsional yang dirumuskan pada tahap ini mencakup login pengguna, pengisian formulir kelayakan multi-langkah, perhitungan SAW otomatis, penyimpanan hasil submission, tampilan detail sekolah, dashboard admin sekolah, dashboard super admin, serta pengelolaan data sekolah dan log aktivitas. Hasil requirements planning menjadi dasar seluruh keputusan desain pada tahap berikutnya.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Tahap User Design", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Tahap user design menerjemahkan kebutuhan pengguna ke dalam struktur antarmuka yang sederhana dan bertahap. Halaman publik dirancang agar pengguna umum dapat memulai dari informasi sekolah, lalu beralih ke menu kelayakan yang dibagi menjadi langkah biodata, domisili, rapor TK, prestasi, serta dokumen dan ekonomi agar proses pengisian tidak terasa berat pada satu halaman panjang.",
            "Rancangan dashboard admin sekolah dan super admin dibuat terpisah karena kebutuhan operasional keduanya berbeda. Admin sekolah difokuskan pada pemantauan profil sekolah, pengajuan kriteria, dan daftar pengguna yang direkomendasikan, sedangkan super admin difokuskan pada monitoring sistem, pengelolaan master sekolah, data pengguna, dan arsip submission pada tingkat keseluruhan.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Pengembangan Sistem", style="Heading 2")
    anchor = insert_paragraph_after(anchor, "Tahap Construction", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Tahap construction dilakukan dengan memisahkan pengembangan frontend, backend, dan basis data agar proses iterasi lebih cepat. Antarmuka pengguna dibangun menggunakan React dan Vite, layanan backend dibangun menggunakan Express, sedangkan penyimpanan data menggunakan MySQL untuk menampung akun, data sekolah, log aktivitas, pengajuan kriteria, dan submission hasil kelayakan.",
            "Logika utama metode SAW ditempatkan pada modul kelayakan sehingga hasil perhitungan dapat tampil secara langsung setelah pengguna menyelesaikan seluruh tahapan input. Pengembangan modul backend memastikan bahwa hasil perhitungan, data sekolah, dan data pengguna dapat dipertukarkan secara konsisten melalui endpoint aplikasi yang terstruktur.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Tahap Cutover", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Tahap cutover pada penelitian ini dilakukan dalam lingkungan implementasi lokal yang telah terhubung dengan basis data MySQL. Sistem memuat akun admin sekolah untuk SD Negeri 1 Kranji dan SD Negeri 1 Sokanegara, akun super admin, serta data awal sekolah sehingga proses uji coba dapat dilakukan tanpa menunggu penyusunan data dari nol.",
            "Hasil cutover internal menunjukkan bahwa sistem telah menghasilkan empat submission riil pada rentang 4 Mei 2026 sampai 5 Mei 2026. Catatan tersebut menandakan bahwa proses login, pengisian form, perhitungan SAW, penyimpanan hasil, dan pembentukan rekomendasi sekolah telah berjalan pada lingkungan pengujian sebelum dipakai untuk evaluasi lapangan yang lebih luas.",
        ],
    )

    anchor = insert_paragraph_after(anchor, "Pengujian Sistem", style="Heading 2")
    anchor = insert_paragraph_after(anchor, "Pengujian Black Box", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Pengujian black box difokuskan pada fitur utama yang berhubungan langsung dengan alur pengguna dan admin sekolah. Pengujian dilakukan dengan melihat kesesuaian keluaran sistem terhadap skenario masukan tanpa menilai struktur kode program di dalamnya.",
            "Ruang lingkup pengujian mencakup autentikasi, pengisian data kelayakan, pembentukan skor otomatis, penyimpanan hasil, dan akses data rekomendasi oleh admin sekolah. Ringkasan hasil pengujian black box disajikan pada Tabel 4.13.",
        ],
    )

    table_413 = [
        ["No", "Skenario Uji", "Hasil yang Diharapkan", "Hasil Aktual", "Status"],
        ["1", "Login admin sekolah", "Admin masuk ke dashboard sesuai akun sekolah.", "Dashboard admin tampil sesuai akun.", "Valid"],
        ["2", "Login pengguna umum dan akses menu kelayakan", "Form kelayakan multi-langkah dapat dibuka.", "Form dapat diakses oleh pengguna.", "Valid"],
        ["3", "Input domisili melalui alamat atau peta", "Kelurahan, kecamatan, dan skor C2 terbentuk otomatis.", "Data wilayah dan skor C2 muncul sesuai input.", "Valid"],
        ["4", "Input rapor dan prestasi", "Sistem mengubah input menjadi skor C3, C4, dan C5.", "Skor terbentuk sesuai data yang diisi.", "Valid"],
        ["5", "Submit proses kelayakan", "Nilai V dan ranking sekolah ditampilkan.", "Hasil rekomendasi sekolah tampil.", "Valid"],
        ["6", "Simpan submission", "Data pengajuan tersimpan ke basis data.", "Data tersimpan pada tabel eligibility_submissions.", "Valid"],
        ["7", "Akses daftar pengguna direkomendasikan oleh admin sekolah", "Admin melihat daftar sesuai sekolah tujuan.", "Daftar pengguna muncul sesuai filter sekolah.", "Valid"],
    ]
    expl_413 = [
        "Tabel 4.13 menunjukkan bahwa seluruh skenario utama menghasilkan keluaran yang sesuai dengan tujuan sistem. Validitas ini penting karena fitur-fitur tersebut merupakan jalur inti yang paling sering digunakan pada saat pengguna melakukan pre-screening dan saat sekolah membaca hasil rekomendasi.",
        "Konsistensi hasil pengujian black box juga didukung oleh keberadaan data submission riil pada basis data. Keberhasilan penyimpanan empat submission internal menjadi indikator bahwa alur dari input pengguna hingga pembacaan data oleh sistem telah berjalan secara utuh dalam satu siklus layanan.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.13 Hasil Pengujian Black Box",
        table_413,
        expl_413,
        font_size=9,
        left_align_cols={1, 2, 3},
    )

    anchor = insert_paragraph_after(anchor, "Rencana User Acceptance Testing (UAT)", style="Heading 3")
    anchor = add_paragraphs(
        anchor,
        [
            "Pengujian UAT pada penelitian ini direncanakan melibatkan dua kelompok responden, yaitu admin sekolah dasar sebagai pihak internal dan pengguna dari TK atau PAUD sebagai pihak eksternal yang mewakili proses pengisian data awal anak. Pemilihan dua kelompok tersebut dilakukan agar evaluasi sistem tidak hanya menilai akurasi hasil, tetapi juga menilai kemudahan penggunaan dari sisi operasional lapangan.",
            "Posisi UAT pada Bab 4 disajikan sebagai rencana pelaksanaan yang siap dijalankan setelah versi antarmuka final disepakati. Bentuk penyajian ini dipilih agar penelitian tetap jujur terhadap status data lapangan, sekaligus menyiapkan instrumen evaluasi yang sesuai dengan kebutuhan sekolah dasar dan mitra TK atau PAUD.",
        ],
    )

    table_414 = [
        ["Kelompok Responden", "Jumlah", "Fokus Penilaian", "Aktivitas Uji", "Keluaran"],
        ["Admin sekolah dasar", "2 orang", "Ketepatan hasil, kemudahan monitoring, dan relevansi data rekomendasi.", "Login admin, cek submission, dan verifikasi daftar pengguna direkomendasikan.", "Skor kepuasan admin dan catatan perbaikan fitur internal."],
        ["Pengguna dari TK/PAUD", "Minimal 3 orang", "Kemudahan input, kejelasan kriteria, dan pemahaman hasil rekomendasi.", "Mengisi form kelayakan end-to-end menggunakan data anak yang didampingi.", "Skor kepuasan user dan masukan perbaikan antarmuka."],
    ]
    expl_414 = [
        "Tabel 4.14 memperlihatkan bahwa fokus penilaian UAT dibedakan berdasarkan peran responden. Admin sekolah dasar diarahkan untuk menilai kegunaan sistem dari sisi pengelolaan dan pemantauan hasil, sedangkan responden dari TK atau PAUD diarahkan untuk menilai kemudahan pengisian data dan keterbacaan rekomendasi yang diberikan sistem.",
        "Pembagian kelompok ini membuat hasil UAT lebih tajam karena setiap responden hanya menilai fungsi yang benar-benar sesuai dengan kebutuhannya. Struktur responden tersebut juga membantu peneliti memisahkan masukan terkait akurasi hasil administrasi dari masukan terkait pengalaman penggunaan form di sisi pengguna umum.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.14 Rencana Pelaksanaan UAT",
        table_414,
        expl_414,
        font_size=9,
        left_align_cols={0, 2, 3, 4},
    )

    table_415 = [
        ["No", "Pernyataan UAT Admin SD", "Aspek"],
        ["1", "Dashboard admin mudah dipahami saat pertama kali digunakan.", "Kemudahan penggunaan"],
        ["2", "Data submission calon siswa tampil jelas dan mudah dibaca.", "Kejelasan informasi"],
        ["3", "Hasil rekomendasi sekolah sesuai dengan kebutuhan seleksi awal.", "Kesesuaian fungsi"],
        ["4", "Detail skor SAW membantu admin menjelaskan alasan rekomendasi.", "Transparansi hasil"],
        ["5", "Sistem membantu mempercepat proses peninjauan data calon siswa.", "Efisiensi kerja"],
    ]
    expl_415 = [
        "Tabel 4.15 memuat kisi-kisi pertanyaan UAT untuk admin sekolah dasar. Setiap pernyataan diarahkan pada kemampuan sistem dalam membantu pekerjaan administratif, mulai dari membaca data submission sampai menafsirkan hasil rekomendasi yang dihasilkan metode SAW.",
        "Arah evaluasi pada tabel ini sengaja difokuskan pada manfaat operasional karena admin sekolah merupakan pihak yang akan memakai sistem sebagai alat bantu seleksi awal. Hasil penilaian dari kelompok admin diharapkan dapat menunjukkan apakah sistem benar-benar mengurangi beban kerja manual yang sebelumnya terjadi.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.15 Kisi-Kisi UAT untuk Admin Sekolah Dasar",
        table_415,
        expl_415,
        font_size=9,
        left_align_cols={1, 2},
    )

    table_416 = [
        ["No", "Pernyataan UAT Pengguna TK/PAUD", "Aspek"],
        ["1", "Form pengisian data calon siswa mudah diikuti dari awal sampai akhir.", "Kemudahan penggunaan"],
        ["2", "Istilah kriteria pada form mudah dipahami oleh pengguna.", "Kejelasan bahasa"],
        ["3", "Proses pengisian domisili, rapor, dan dokumen terasa runtut.", "Alur interaksi"],
        ["4", "Hasil rekomendasi sekolah mudah dipahami setelah perhitungan selesai.", "Kejelasan hasil"],
        ["5", "Sistem bermanfaat sebagai simulasi awal sebelum mendaftar ke SD.", "Kemanfaatan sistem"],
    ]
    expl_416 = [
        "Tabel 4.16 memuat kisi-kisi pertanyaan UAT untuk pengguna dari TK atau PAUD yang berperan sebagai pendamping pengisian data awal anak. Isi pertanyaan dirancang untuk menilai apakah alur form, istilah yang dipakai, dan tampilan hasil sudah cukup mudah dipahami oleh pengguna non-teknis.",
        "Kelompok pengguna ini penting karena mereka mewakili kondisi lapangan ketika orang tua atau pendamping pendidikan anak usia dini mencoba melakukan simulasi kelayakan secara mandiri. Umpan balik dari kelompok ini diharapkan dapat membantu penyempurnaan antarmuka sebelum sistem dipakai lebih luas.",
    ]
    anchor = add_table_section(
        anchor,
        "Tabel 4.16 Kisi-Kisi UAT untuk Pengguna dari TK/PAUD",
        table_416,
        expl_416,
        font_size=9,
        left_align_cols={1, 2},
    )

    anchor = add_paragraphs(
        anchor,
        [
            "Pengolahan nilai UAT nantinya dilakukan menggunakan skala Likert 1 sampai 5 dengan rumus persentase kelayakan, yaitu total skor yang diperoleh dibagi skor ideal kemudian dikalikan 100 persen. Mekanisme ini dipilih karena mudah dipahami dan sejalan dengan formula evaluasi penerimaan sistem yang telah dijelaskan pada Bab II.",
            "Nilai numerik UAT belum disajikan pada tahap ini karena pengujian lapangan kepada admin sekolah dasar dan responden dari TK atau PAUD direncanakan dilakukan setelah tahap revisi antarmuka final selesai. Kondisi tersebut membuat Bab 4 tetap akurat secara akademik sekaligus siap dipakai sebagai dasar pelaksanaan UAT pada tahap berikutnya.",
        ],
    )

    doc.save(str(OUTPUT_DOC))
    shutil.copyfile(OUTPUT_DOC, FINAL_DOC)


if __name__ == "__main__":
    build_document()
