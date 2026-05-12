from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


WORKDIR = Path(__file__).resolve().parent.parent
INPUT_DOC = WORKDIR / "working-bab4-revisi.docx"
OUTPUT_DOC = WORKDIR / "working-bab4-final-rad-saw.docx"
ASSET_DIR = WORKDIR / "docx_assets"
ASSET_DIR.mkdir(exist_ok=True)


def get_font(size: int, bold: bool = False):
    candidates = []
    if bold:
        candidates.extend(
            [
                r"C:\Windows\Fonts\arialbd.ttf",
                r"C:\Windows\Fonts\calibrib.ttf",
            ]
        )
    candidates.extend(
        [
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
        ]
    )
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = []
    for word in words:
        trial = " ".join(current + [word]).strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current.append(word)
        else:
            if current:
                lines.append(" ".join(current))
            current = [word]
    if current:
        lines.append(" ".join(current))
    return lines


def draw_box(draw, xy, text, fill, outline, text_color=(20, 20, 20), font=None, radius=20):
    font = font or get_font(28)
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=3)
    max_width = int((x2 - x1) - 40)
    lines = wrap_text(draw, text, font, max_width)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 6
    total_height = len(lines) * line_height
    y = y1 + ((y2 - y1) - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=text_color)
        y += line_height


def draw_arrow(draw, start, end, fill=(74, 101, 148), width=5):
    draw.line([start, end], fill=fill, width=width)
    ex, ey = end
    sx, sy = start
    if ex != sx:
        direction = 1 if ex > sx else -1
        draw.polygon(
            [
                (ex, ey),
                (ex - 16 * direction, ey - 10),
                (ex - 16 * direction, ey + 10),
            ],
            fill=fill,
        )
    else:
        direction = 1 if ey > sy else -1
        draw.polygon(
            [
                (ex, ey),
                (ex - 10, ey - 16 * direction),
                (ex + 10, ey - 16 * direction),
            ],
            fill=fill,
        )


def create_rad_saw_diagram(path: Path):
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(38, bold=True)
    box_font = get_font(26, bold=True)
    body_font = get_font(24)

    draw.text((70, 40), "Integrasi Metode RAD dan SAW pada Penelitian", fill=(18, 42, 76), font=title_font)
    draw.text(
        (70, 100),
        "RAD membungkus seluruh siklus pengembangan sistem, sedangkan SAW dipakai pada modul inti penilaian kelayakan.",
        fill=(60, 60, 60),
        font=body_font,
    )

    boxes = [
        ((80, 220, 430, 390), "Requirements Planning", (227, 239, 255)),
        ((500, 220, 850, 390), "User Design", (224, 242, 241)),
        ((920, 220, 1270, 390), "Construction", (255, 238, 214)),
        ((1340, 220, 1690, 390), "Cutover", (234, 229, 255)),
    ]
    for xy, text, fill in boxes:
        draw_box(draw, xy, text, fill, (94, 121, 167), font=box_font)

    draw_arrow(draw, (430, 305), (500, 305))
    draw_arrow(draw, (850, 305), (920, 305))
    draw_arrow(draw, (1270, 305), (1340, 305))

    saw_outer = (420, 520, 1360, 820)
    draw.rounded_rectangle(saw_outer, radius=26, outline=(35, 94, 78), fill=(239, 250, 246), width=4)
    draw.text((455, 550), "Modul SAW pada Sistem Kelayakan", fill=(24, 73, 62), font=title_font)

    inner_boxes = [
        ((470, 640, 700, 760), "Input skor\nkriteria", (255, 255, 255)),
        ((745, 640, 975, 760), "Normalisasi\nmatriks", (255, 255, 255)),
        ((1020, 640, 1250, 760), "Nilai preferensi\n(V)", (255, 255, 255)),
    ]
    for xy, text, fill in inner_boxes:
        draw_box(draw, xy, text, fill, (35, 94, 78), font=box_font, radius=18)

    draw_arrow(draw, (700, 700), (745, 700), fill=(35, 94, 78))
    draw_arrow(draw, (975, 700), (1020, 700), fill=(35, 94, 78))

    draw_box(
        draw,
        (1295, 620, 1620, 780),
        "Output:\nranking sekolah\n& rekomendasi",
        (255, 245, 245),
        (176, 69, 69),
        font=box_font,
        radius=18,
    )
    draw_arrow(draw, (1250, 700), (1295, 700), fill=(176, 69, 69))

    image.save(path)


def create_architecture_diagram(path: Path):
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    title_font = get_font(38, bold=True)
    box_font = get_font(26, bold=True)
    body_font = get_font(24)

    draw.text((70, 40), "Arsitektur Implementasi Sistem DILAYAKIN", fill=(18, 42, 76), font=title_font)
    draw.text(
        (70, 100),
        "Aplikasi memakai frontend React, backend Node.js + Express, dan basis data MySQL untuk menyimpan hasil simulasi kelayakan.",
        fill=(60, 60, 60),
        font=body_font,
    )

    actors = [
        ((120, 250, 420, 420), "Pengguna umum\n(orang tua / pendamping TK-PAUD)", (240, 246, 255)),
        ((120, 470, 420, 640), "Admin sekolah dasar", (240, 255, 244)),
        ((120, 690, 420, 860), "Super admin", (255, 248, 235)),
    ]
    for xy, text, fill in actors:
        draw_box(draw, xy, text, fill, (100, 114, 144), font=box_font)

    draw_box(
        draw,
        (560, 320, 980, 790),
        "Frontend\nReact.js + Vite + Tailwind CSS\n\nForm kelayakan multi-langkah\nDashboard admin sekolah\nDashboard super admin",
        (247, 250, 255),
        (74, 101, 148),
        font=box_font,
    )
    draw_box(
        draw,
        (1090, 320, 1510, 790),
        "Backend API\nNode.js + Express.js\n\nAutentikasi\nPerhitungan SAW\nKriteria, log aktivitas,\nrekomendasi dan ekspor data",
        (250, 247, 255),
        (106, 76, 147),
        font=box_font,
    )
    draw_box(
        draw,
        (1560, 420, 1730, 690),
        "MySQL\n\nusers\nschools\neligibility_submissions\ncriteria_requests\nactivity_logs",
        (255, 250, 240),
        (148, 95, 52),
        font=box_font,
    )

    for y in (335, 555, 775):
        draw_arrow(draw, (420, y), (560, 555), fill=(74, 101, 148))
    draw_arrow(draw, (980, 555), (1090, 555), fill=(106, 76, 147))
    draw_arrow(draw, (1510, 555), (1560, 555), fill=(148, 95, 52))

    image.save(path)


def set_text(paragraph: Paragraph, text: str):
    while paragraph.runs:
        paragraph._p.remove(paragraph.runs[0]._r)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        set_text(new_paragraph, text)
    return new_paragraph


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        set_text(new_paragraph, text)
    return new_paragraph


def insert_picture_after(paragraph: Paragraph, image_path: Path, width_inches: float = 6.3) -> Paragraph:
    new_paragraph = insert_paragraph_after(paragraph)
    new_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new_paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    return new_paragraph


def apply_reference_format(paragraph: Paragraph):
    paragraph.style = "Normal"
    paragraph.paragraph_format.left_indent = Cm(0.85)
    paragraph.paragraph_format.first_line_indent = Cm(-0.85)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def find_paragraph(document: Document, text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def main():
    rad_path = ASSET_DIR / "bab4-rad-saw.png"
    arch_path = ASSET_DIR / "bab4-arsitektur.png"
    create_rad_saw_diagram(rad_path)
    create_architecture_diagram(arch_path)

    document = Document(str(INPUT_DOC))

    # Bab II: sesuaikan teori dengan stack project aktual.
    p = find_paragraph(document, "Hypertext Preprocessor (PHP)")
    set_text(p, "React.js")
    p = find_paragraph(document, "Menurut buku yang ditulis oleh Tim EMS, Hypertext Preprocessor (PHP) dipahami sebagai bahasa scripting yang disisipkan dalam Hypertext Markup Language (HTML) untuk membangun halaman web dinamis. PHP bekerja pada sisi server, mampu menghasilkan tampilan yang interaktif, serta menjadi salah satu bahasa pemrograman yang paling populer karena dukungan berbagai web server dan kemudahan instalasinya. PHP juga dianggap mudah dipelajari karena memiliki banyak referensi dan komunitas pengembang, serta bersifat open source sehingga dapat dijalankan di berbagai sistem operasi dan digunakan untuk membangun berbagai jenis aplikasi web .")
    set_text(
        p,
        "Menurut Chen, Thaduri, dan Ballamudi, React.js merupakan library JavaScript yang berfokus pada pembentukan antarmuka pengguna berbasis komponen, virtual DOM, dan pemrosesan tampilan yang responsif. Pendekatan tersebut membuat pengembangan halaman web menjadi lebih terstruktur karena setiap bagian antarmuka dapat dirancang sebagai komponen yang dapat digunakan kembali pada halaman lain dengan alur data yang tetap jelas. Penjelasan serupa juga ditegaskan oleh Kulkarni dkk. yang menempatkan React.js sebagai komponen frontend utama pada pengembangan full-stack modern karena mampu meningkatkan responsivitas tampilan dan mendukung pengembangan aplikasi web yang dinamis.",
    )
    p = find_paragraph(document, "Pemahaman terhadap definisi PHP dari referensi tersebut memperlihatkan bahwa PHP merupakan bahasa pemrograman yang efektif untuk membangun aplikasi web dinamis, efisien, dan mudah dikembangkan. Dalam konteks penelitian yang sedang dilakukan, PHP menjadi teknologi yang relevan karena mampu mendukung proses perancangan sistem kelayakan penerimaan siswa yang membutuhkan pengelolaan data secara real-time. PHP juga mendukung integrasi dengan basis data dan memungkinkan penyajian informasi secara cepat, sehingga sesuai untuk diimplementasikan dalam sistem informasi berbasis web yang akan dikembangkan.")
    set_text(
        p,
        "Pemahaman tersebut menunjukkan bahwa React.js relevan digunakan pada penelitian ini karena sistem kelayakan calon siswa membutuhkan form multi-langkah, tampilan hasil simulasi yang interaktif, serta dashboard admin yang dapat menampilkan data secara cepat. Penerapan React.js pada project ini membantu peneliti membangun antarmuka yang modular, mudah dipelihara, dan sesuai dengan kebutuhan pengguna umum, admin sekolah dasar, dan super admin pada lingkungan aplikasi berbasis web.",
    )

    p = find_paragraph(document, "Laravel")
    set_text(p, "Node.js dan Express.js")
    p = find_paragraph(document, "Menurut Fauzi dan Darmawan, Laravel merupakan framework pengembangan website berbasis PHP yang membantu proses pembangunan aplikasi menjadi lebih rapi, cepat, dan aman karena menyediakan berbagai fitur modern seperti routing, Object-Relational Mapping (ORM), serta dukungan komunitas yang luas . Menurut Yansyah, Auliana, dan Setya, Laravel adalah framework PHP yang menawarkan struktur Model-View-Controller (MVC) yang rapi, fitur keamanan, validasi data, manajemen basis data yang mudah, serta kemudahan dalam membangun aplikasi web yang kompleks, sehingga sangat sesuai untuk pengembangan aplikasi akademik yang membutuhkan pengelolaan data terstruktur dan akses informasi secara real-time .")
    set_text(
        p,
        "Menurut Jadhav dan Gonsalves, Node.js merupakan runtime JavaScript berbasis event-driven dan non-blocking I/O yang dirancang untuk membangun aplikasi jaringan dan layanan web secara efisien. Kulkarni dkk. menambahkan bahwa Express.js berperan sebagai framework backend yang menyederhanakan pembuatan routing, pengelolaan request-response, dan pembentukan REST API sehingga proses komunikasi antara frontend dan basis data dapat berjalan lebih ringkas serta terstruktur.",
    )
    p = find_paragraph(document, "Berdasarkan kedua referensi tersebut, dapat disimpulkan bahwa Laravel merupakan framework yang dirancang untuk mempercepat, merapikan, serta mengamankan proses pengembangan aplikasi web dengan menyediakan komponen siap pakai dan arsitektur MVC yang jelas. Hubungannya dengan penelitian ini, pemilihan Laravel menjadi relevan karena kebutuhan sistem kelayakan calon siswa memerlukan pengelolaan data yang terstruktur, proses validasi yang jelas, serta tampilan antarmuka yang mudah dikembangkan. Dengan memanfaatkan fitur bawaan Laravel, sistem dapat dibangun lebih efisien dan mampu mendukung proses seleksi yang akurat, terstandar, serta mudah dikembangkan pada tahap berikutnya.")
    set_text(
        p,
        "Berdasarkan kedua referensi tersebut, Node.js dan Express.js dapat dipahami sebagai kombinasi teknologi backend yang mendukung pembangunan aplikasi web secara cepat, efisien, dan mudah dikembangkan. Relevansinya terhadap penelitian ini sangat kuat karena sistem kelayakan calon siswa memerlukan layanan autentikasi, penyimpanan submission, pengelolaan data sekolah, log aktivitas, dan pemrosesan hasil SAW yang harus dapat diakses oleh antarmuka React.js secara konsisten melalui endpoint aplikasi.",
    )

    p = find_paragraph(document, "Berdasarkan kedua pengertian tersebut, dapat disimpulkan bahwa MySQL merupakan sistem manajemen basis data relasional yang efisien, fleksibel, dan mudah diimplementasikan pada berbagai platform. Dalam konteks penelitian sistem kelayakan calon siswa, penggunaan MySQL menjadi relevan karena mampu mendukung kebutuhan penyimpanan dan pengelolaan data secara terstruktur, stabil, dan mudah diintegrasikan dengan bahasa pemrograman web seperti PHP. Hal ini membantu memastikan bahwa sistem yang dikembangkan dapat berjalan secara optimal, konsisten, dan dapat menangani kebutuhan data aplikasi secara efektif.")
    set_text(
        p,
        "Berdasarkan kedua pengertian tersebut, dapat disimpulkan bahwa MySQL merupakan sistem manajemen basis data relasional yang efisien, fleksibel, dan mudah diimplementasikan pada berbagai platform. Dalam konteks penelitian sistem kelayakan calon siswa, penggunaan MySQL menjadi relevan karena mampu mendukung kebutuhan penyimpanan dan pengelolaan data secara terstruktur, stabil, dan mudah diintegrasikan dengan layanan backend berbasis Node.js dan Express.js. Kondisi tersebut membantu memastikan bahwa sistem yang dikembangkan dapat berjalan secara optimal, konsisten, dan mampu menangani kebutuhan data aplikasi secara efektif.",
    )

    p = find_paragraph(document, "Meskipun Aldi Ramadani hanya menyebutkan diagram ini sebagai salah satu alat desain , menurut referensi yang digunakan oleh Nauval Dwi Rivalfakhri dan Apriade Voutama, Activity Diagram digunakan untuk memodelkan alur proses atau aktivitas yang terjadi dalam sistem . Referensi lain dalam jurnal yang sama menegaskan bahwa diagram ini memvisualisasikan urutan langkah-langkah yang diambil oleh pengguna atau sistem dalam menjalankan aktivitas tertentu.")
    set_text(
        p,
        "Activity Diagram digunakan untuk memodelkan alur proses atau aktivitas yang terjadi dalam sistem sebagaimana dijelaskan oleh referensi yang digunakan oleh Nauval Dwi Rivalfakhri dan Apriade Voutama. Referensi lain dalam jurnal yang sama menegaskan bahwa diagram ini memvisualisasikan urutan langkah-langkah yang diambil oleh pengguna atau sistem dalam menjalankan aktivitas tertentu.",
    )

    # Bab III: perbarui alat, tahapan, dan target pengujian.
    table = document.tables[7]
    rows = [
        ("1.", "Hardware (Perangkat Keras)", "Laptop ASUS VivoBook", "Processor: AMD Ryzen 5 5500U, RAM: 16 GB, Storage: SSD 477 GB NVMe, GPU: AMD Radeon Graphics."),
        ("2.", "Sistem Operasi", "Windows 11", "Sistem operasi 64-bit yang digunakan sebagai lingkungan pengembangan, pengujian, dan render antarmuka aplikasi."),
        ("3.", "Frontend", "JavaScript, React.js, Vite, Tailwind CSS", "Digunakan untuk membangun antarmuka multi-langkah, dashboard admin, serta tampilan hasil simulasi kelayakan secara responsif."),
        ("4.", "Backend", "Node.js dan Express.js", "Digunakan untuk membangun layanan autentikasi, pengelolaan data sekolah, penyimpanan submission, log aktivitas, dan integrasi proses SAW."),
        ("5.", "Basis Data", "MySQL", "Digunakan untuk menyimpan data akun, data sekolah, hasil submission kelayakan, pengajuan kriteria, dan log aktivitas sistem."),
        ("6.", "IDE & Modeling", "Visual Studio Code, Figma, StarUML", "Digunakan untuk penulisan kode, desain antarmuka, serta pemodelan diagram UML dan struktur sistem."),
        ("7.", "Pengujian", "Google Chrome dan Postman", "Digunakan untuk pengujian fungsional antarmuka web, validasi endpoint API, dan verifikasi hasil proses sistem."),
    ]
    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = value

    p = find_paragraph(document, "Diagram alur penelitian menyajikan rangkaian tahapan kegiatan penelitian yang disusun secara terstruktur dan berurutan, dimulai dari proses identifikasi permasalahan, kajian literatur, hingga tahap penarikan kesimpulan. Penyusunan diagram ini mengacu pada siklus pengembangan perangkat lunak menggunakan metode Rapid Application Develompent (RAD).")
    set_text(
        p,
        "Diagram alur penelitian menyajikan rangkaian tahapan kegiatan penelitian yang disusun secara terstruktur dan berurutan, dimulai dari proses identifikasi permasalahan, kajian literatur, hingga tahap penarikan kesimpulan. Penyusunan diagram ini mengacu pada siklus pengembangan perangkat lunak menggunakan metode Rapid Application Development (RAD).",
    )

    p = find_paragraph(document, "Tahap pengumpulan data merupakan proses awal penelitian yang bertujuan untuk memperoleh informasi yang relevan dan akurat sebagai dasar perancangan sistem. Data dikumpulkan melalui observasi wawancara dengan panitia PPDB langsung terhadap proses PPDB di SD Negeri 1 Sokanegara dan SD Negeri 1 Kranji serta terkait mekanisme pendaftaran, penentuan kelayakan calon siswa, dan kendala yang dihadapi. Selain itu, dilakukan pengumpulan data sekunder berupa dokumen kebijakan PPDB, regulasi usia, serta referensi penelitian terdahulu yang berkaitan dengan Rapid Application Develompent (RAD) dan metode Simple Additive Weighting (SAW).")
    set_text(
        p,
        "Tahap pengumpulan data merupakan proses awal penelitian yang bertujuan untuk memperoleh informasi yang relevan dan akurat sebagai dasar perancangan sistem. Data dikumpulkan melalui observasi dan wawancara langsung dengan panitia PPDB di SD Negeri 1 Sokanegara dan SD Negeri 1 Kranji mengenai mekanisme pendaftaran, penentuan kelayakan calon siswa, serta kendala yang dihadapi. Pengumpulan data sekunder dilakukan melalui dokumen kebijakan PPDB, regulasi usia, statistik pendidikan, dan referensi penelitian terdahulu yang berkaitan dengan Rapid Application Development (RAD) serta metode Simple Additive Weighting (SAW).",
    )

    p = find_paragraph(document, "Tahap ini merupakan bagian dari fase perencanaan kebutuhan pada metode Rapid Application Develompent. Analisis dilakukan untuk menentukan kebutuhan pengguna dan kebutuhan sistem. Pengguna sistem adalah calon siswa atau wali murid. Kebutuhan sistem mencakup pengelolaan data kriteria, input data calon siswa, proses perhitungan kelayakan menggunakan metode SAW, serta penyajian hasil evaluasi secara jelas. Kebutuhan tersebut kemudian dirumuskan dalam bentuk fitur-fitur utama yang akan dikembangkan secara bertahap.")
    set_text(
        p,
        "Tahap ini merupakan bagian dari fase perencanaan kebutuhan pada metode Rapid Application Development. Analisis dilakukan untuk menentukan kebutuhan pengguna dan kebutuhan sistem. Pengguna sistem mencakup pengguna umum, admin sekolah dasar, dan super admin. Kebutuhan sistem meliputi pengelolaan data sekolah, input data calon siswa, proses perhitungan kelayakan menggunakan metode SAW, penyimpanan submission, pengajuan perubahan kriteria, dan penyajian hasil evaluasi secara jelas. Kebutuhan tersebut kemudian dirumuskan dalam bentuk fitur-fitur utama yang akan dikembangkan secara bertahap.",
    )

    p = find_paragraph(document, "Proses visualisasi antarmuka pengguna dimulai melalui pembuatan rancangan Low-Fidelity (Lo-Fi). Pemetaan struktur serta tata letak informasi dilakukan untuk mematangkan persiapan sebelum fase pengodean dan implementasi algoritma SAW dimulai. Visualisasi rancangan yang sederhana bertujuan membantu pengguna dalam memahami mekanisme kerja sistem secara cepat.")
    set_text(
        p,
        "Proses visualisasi antarmuka pengguna dimulai melalui pembuatan rancangan low-fidelity untuk halaman publik, form kelayakan, dashboard admin sekolah, dan dashboard super admin. Pemetaan struktur serta tata letak informasi dilakukan untuk mematangkan persiapan sebelum fase pengodean dan implementasi algoritma SAW dimulai. Visualisasi rancangan yang sederhana membantu pengguna memahami mekanisme kerja sistem secara cepat.",
    )

    p = find_paragraph(document, "Menu Dashboard dan fitur Pre-Screening dapat diakses langsung melalui bilah navigasi yang tersedia. Halaman utama menyediakan formulir input data kriteria yang harus diisi oleh calon pendaftar. Variabel input tersebut merupakan wujud visual dari kriteria C1 dan C2 yang diolah menggunakan rumus Simple Additive Weighting. Kelancaran transisi menuju tahap konstruksi kode dijamin oleh dokumentasi visual ini agar hasil akhir sistem sesuai dengan ekspektasi operasional sekolah.")
    set_text(
        p,
        "Menu dashboard, detail sekolah, dan fitur pre-screening dapat diakses langsung melalui bilah navigasi yang tersedia. Halaman utama menyediakan formulir multi-langkah yang memuat biodata, domisili, rapor TK, prestasi, dokumen, dan kondisi ekonomi. Setiap elemen input tersebut dirancang agar sesuai dengan kriteria SAW yang diolah sistem, sehingga dokumentasi visual ini menjadi dasar penting agar hasil implementasi sesuai dengan ekspektasi operasional sekolah.",
    )

    p = find_paragraph(document, "Tahap konstruksi melibatkan proses transformasi hasil desain ke dalam baris kode program yang fungsional. Pengembangan sistem dilakukan secara cepat menggunakan framework Laravel untuk mengimplementasikan logika algoritma SAW. Proses pengembangan difokuskan pada integrasi modul perhitungan skor preferensi, pengelolaan data kriteria, serta sinkronisasi antarmuka pengguna. Validasi mandiri dilakukan secara berkala selama proses pengodean guna memastikan keakuratan hasil penilaian kelayakan bagi setiap calon siswa.")
    set_text(
        p,
        "Tahap konstruksi melibatkan proses transformasi hasil desain ke dalam baris kode program yang fungsional. Pengembangan sistem dilakukan dengan memisahkan frontend React.js, backend Node.js + Express.js, dan basis data MySQL agar proses iterasi lebih terarah. Proses pengembangan difokuskan pada integrasi modul perhitungan skor preferensi SAW, pengelolaan data sekolah, penyimpanan submission, serta sinkronisasi antarmuka pengguna. Validasi mandiri dilakukan secara berkala selama proses pengodean guna memastikan keakuratan hasil penilaian kelayakan bagi setiap calon siswa.",
    )

    p = find_paragraph(document, "Tahap ini dilakukan untuk memastikan bahwa sistem telah memenuhi kebutuhan dan harapan pengguna akhir. Pengujian dilakukan dengan melibatkan pihak yang membutuhkan sistem, khususnya calon siswa atau wali murid. Apabila hasil pengujian telah sesuai, maka sistem dinyatakan dapat diterima dan siap digunakan. Umpan balik pengguna menjadi dasar validasi akhir sistem.")
    set_text(
        p,
        "Tahap ini dilakukan untuk memastikan bahwa sistem telah memenuhi kebutuhan dan harapan pengguna akhir. Pengujian direncanakan melibatkan dua kelompok, yaitu admin sekolah dasar sebagai pengguna internal dan pendamping dari TK atau PAUD sebagai pengguna eksternal yang mewakili pengisian data awal anak. Apabila hasil pengujian telah sesuai, sistem dinyatakan dapat diterima dan siap digunakan. Umpan balik dari kedua kelompok tersebut menjadi dasar validasi akhir sistem.",
    )

    # Bab IV: finalisasi struktur dan akurasi isi.
    p = find_paragraph(document, "Pengumpulan dan Analisis Data")
    set_text(p, "4.1 Hasil")
    p.style = document.styles["Heading 2"]

    p = find_paragraph(document, "Data aplikasi memperlihatkan bahwa sistem yang dibangun telah menyimpan empat submission riil pada tabel eligibility_submissions sampai tanggal 5 Mei 2026. Data tersebut menjadi sumber utama pembahasan pada Bab 4 karena dapat menunjukkan hubungan langsung antara kebutuhan lapangan, implementasi logika SAW, dan hasil rekomendasi sekolah yang dihasilkan sistem.")
    set_text(
        p,
        "Data aplikasi memperlihatkan bahwa sistem yang dibangun telah menyimpan tujuh submission riil pada tabel eligibility_submissions sampai 6 Mei 2026. Data tersebut menjadi sumber utama pembahasan pada Bab 4 karena dapat menunjukkan hubungan langsung antara kebutuhan lapangan, implementasi logika SAW, dan hasil rekomendasi sekolah yang dihasilkan sistem.",
    )

    intro_anchor = find_paragraph(document, "Kriteria yang digunakan dalam perhitungan SAW terdiri atas usia, domisili, nilai rapor TK, prestasi akademik, prestasi non-akademik, kelengkapan dokumen, dan kondisi ekonomi keluarga. Keterlibatan orang tua tetap dikumpulkan sebagai data pendamping, tetapi variabel tersebut tidak dimasukkan ke dalam penjumlahan terbobot agar proses penilaian tetap fokus pada kriteria yang dapat diskor secara konsisten.")
    pic = insert_picture_after(intro_anchor, rad_path, width_inches=6.5)
    caption = insert_paragraph_after(pic, "Gambar 4.1 Integrasi Metode RAD dan SAW pada Penelitian", "Caption")
    expl1 = insert_paragraph_after(
        caption,
        "Gambar 4.1 menunjukkan bahwa metode Rapid Application Development berperan sebagai kerangka besar yang membungkus seluruh kegiatan penelitian, mulai dari perencanaan kebutuhan sampai tahap cutover. Posisi metode SAW berada di dalam modul inti sistem karena metode tersebut digunakan secara spesifik untuk mengubah data calon siswa menjadi nilai preferensi dan rekomendasi sekolah.",
        "Normal",
    )
    insert_paragraph_after(
        expl1,
        "Struktur tersebut menegaskan perbedaan peran antara RAD dan SAW. RAD digunakan untuk mengatur proses pengembangan perangkat lunak secara menyeluruh, sedangkan SAW digunakan untuk menjalankan fungsi sistem pendukung keputusan yang sebelumnya pada referensi pembanding masih menggunakan metode AHP.",
        "Normal",
    )

    p = find_paragraph(document, "Analisis Sistem")
    set_text(p, "4.1.1 Hasil Tahap Requirements Planning")
    p.style = document.styles["Heading 3"]

    p = find_paragraph(document, "Perhitungan Metode Simple Additive Weighting (SAW)")
    set_text(p, "4.1.2 Hasil Implementasi Metode Simple Additive Weighting (SAW)")
    p.style = document.styles["Heading 3"]

    p = find_paragraph(document, "Perancangan Sistem")
    set_text(p, "4.1.3 Hasil Tahap User Design")
    p.style = document.styles["Heading 3"]
    p = find_paragraph(document, "Tahap Requirements Planning")
    set_text(p, "Struktur Navigasi dan Peran Pengguna")
    p.style = document.styles["Heading 4"]
    p = find_paragraph(document, "Tahap requirements planning pada metode RAD dimulai dari pengumpulan kebutuhan pengguna melalui wawancara dengan panitia sekolah dan telaah terhadap alur PPDB yang sudah berjalan. Hasil tahap ini menunjukkan bahwa sistem harus mampu menampung dua kelompok pengguna utama, yaitu pengguna umum yang melakukan pre-screening dan admin sekolah yang memantau hasil rekomendasi.")
    set_text(
        p,
        "Tahap user design menerjemahkan kebutuhan sistem ke dalam struktur navigasi yang mudah dipahami oleh tiga aktor utama, yaitu pengguna umum, admin sekolah dasar, dan super admin. Pengguna umum diarahkan ke halaman informasi sekolah dan form simulasi kelayakan, admin sekolah diarahkan ke dashboard rekomendasi dan pengajuan kriteria, sedangkan super admin diarahkan ke pengelolaan sekolah, pengguna, dan aktivitas sistem.",
    )
    p = find_paragraph(document, "Kebutuhan fungsional yang dirumuskan pada tahap ini mencakup login pengguna, pengisian formulir kelayakan multi-langkah, perhitungan SAW otomatis, penyimpanan hasil submission, tampilan detail sekolah, dashboard admin sekolah, dashboard super admin, serta pengelolaan data sekolah dan log aktivitas. Hasil requirements planning menjadi dasar seluruh keputusan desain pada tahap berikutnya.")
    set_text(
        p,
        "Form kelayakan dirancang dalam beberapa langkah agar pengguna tidak mengisi seluruh data pada satu halaman yang terlalu panjang. Pendekatan bertahap tersebut memudahkan pengguna saat mengisi biodata, domisili, rapor TK, prestasi, dokumen, dan kondisi ekonomi, sekaligus memudahkan sistem untuk membentuk skor kriteria secara otomatis sebelum proses SAW dijalankan.",
    )
    p = find_paragraph(document, "Tahap User Design")
    set_text(p, "Rancangan Antarmuka")
    p.style = document.styles["Heading 4"]
    p = find_paragraph(document, "Tahap user design menerjemahkan kebutuhan pengguna ke dalam struktur antarmuka yang sederhana dan bertahap. Halaman publik dirancang agar pengguna umum dapat memulai dari informasi sekolah, lalu beralih ke menu kelayakan yang dibagi menjadi langkah biodata, domisili, rapor TK, prestasi, serta dokumen dan ekonomi agar proses pengisian tidak terasa berat pada satu halaman panjang.")
    set_text(
        p,
        "Tahap user design menerjemahkan kebutuhan pengguna ke dalam rancangan antarmuka yang sederhana dan bertahap. Halaman publik dirancang agar pengguna umum dapat memulai dari informasi sekolah, lalu beralih ke menu kelayakan yang dibagi menjadi langkah biodata, domisili, rapor TK, prestasi, serta dokumen dan ekonomi agar proses pengisian tidak terasa berat pada satu halaman panjang.",
    )
    p = find_paragraph(document, "Rancangan dashboard admin sekolah dan super admin dibuat terpisah karena kebutuhan operasional keduanya berbeda. Admin sekolah difokuskan pada pemantauan profil sekolah, pengajuan kriteria, dan daftar pengguna yang direkomendasikan, sedangkan super admin difokuskan pada monitoring sistem, pengelolaan master sekolah, data pengguna, dan arsip submission pada tingkat keseluruhan.")
    set_text(
        p,
        "Rancangan dashboard admin sekolah dan super admin dibuat terpisah karena kebutuhan operasional keduanya berbeda. Admin sekolah difokuskan pada pemantauan profil sekolah, pengajuan kriteria, dan daftar pengguna yang direkomendasikan, sedangkan super admin difokuskan pada monitoring sistem, pengelolaan master sekolah, data pengguna, dan arsip submission pada tingkat keseluruhan.",
    )

    p = find_paragraph(document, "Pengembangan Sistem")
    set_text(p, "4.1.4 Hasil Tahap Construction dan Cutover")
    p.style = document.styles["Heading 3"]
    p = find_paragraph(document, "Tahap Construction")
    p.style = document.styles["Heading 4"]
    p = find_paragraph(document, "Tahap Cutover")
    p.style = document.styles["Heading 4"]

    arch_anchor = find_paragraph(document, "Logika utama metode SAW ditempatkan pada modul kelayakan sehingga hasil perhitungan dapat tampil secara langsung setelah pengguna menyelesaikan seluruh tahapan input. Pengembangan modul backend memastikan bahwa hasil perhitungan, data sekolah, dan data pengguna dapat dipertukarkan secara konsisten melalui endpoint aplikasi yang terstruktur.")
    pic = insert_picture_after(arch_anchor, arch_path, width_inches=6.5)
    caption = insert_paragraph_after(pic, "Gambar 4.2 Arsitektur Implementasi Sistem DILAYAKIN", "Caption")
    expl1 = insert_paragraph_after(
        caption,
        "Gambar 4.2 menunjukkan arsitektur implementasi yang dipakai pada project. Seluruh interaksi pengguna diproses melalui antarmuka React.js, kemudian diteruskan ke backend Node.js dan Express.js untuk autentikasi, pengelolaan data, serta pemrosesan hasil SAW sebelum disimpan ke basis data MySQL.",
        "Normal",
    )
    insert_paragraph_after(
        expl1,
        "Pemisahan lapisan frontend, backend, dan basis data membuat proses iterasi pengembangan menjadi lebih cepat sekaligus menjaga struktur aplikasi tetap rapi. Pola tersebut juga memudahkan pengembangan lanjutan karena fitur publik, dashboard admin sekolah, dan dashboard super admin dapat dikembangkan tanpa mengubah logika inti perhitungan secara langsung.",
        "Normal",
    )

    p = find_paragraph(document, "Hasil cutover internal menunjukkan bahwa sistem telah menghasilkan empat submission riil pada rentang 4 Mei 2026 sampai 5 Mei 2026. Catatan tersebut menandakan bahwa proses login, pengisian form, perhitungan SAW, penyimpanan hasil, dan pembentukan rekomendasi sekolah telah berjalan pada lingkungan pengujian sebelum dipakai untuk evaluasi lapangan yang lebih luas.")
    set_text(
        p,
        "Hasil cutover internal menunjukkan bahwa sistem telah menghasilkan tujuh submission riil pada rentang 4 Mei 2026 sampai 6 Mei 2026. Catatan tersebut menandakan bahwa proses login, pengisian form, perhitungan SAW, penyimpanan hasil, dan pembentukan rekomendasi sekolah telah berjalan pada lingkungan pengujian sebelum dipakai untuk evaluasi lapangan yang lebih luas.",
    )

    p = find_paragraph(document, "Pengujian Sistem")
    set_text(p, "4.1.5 Hasil Pengujian Sistem")
    p.style = document.styles["Heading 3"]
    p = find_paragraph(document, "Pengujian Black Box")
    p.style = document.styles["Heading 4"]
    p = find_paragraph(document, "Rencana User Acceptance Testing (UAT)")
    p.style = document.styles["Heading 4"]
    p = find_paragraph(document, "Konsistensi hasil pengujian black box juga didukung oleh keberadaan data submission riil pada basis data. Keberhasilan penyimpanan empat submission internal menjadi indikator bahwa alur dari input pengguna hingga pembacaan data oleh sistem telah berjalan secara utuh dalam satu siklus layanan.")
    set_text(
        p,
        "Konsistensi hasil pengujian black box juga didukung oleh keberadaan data submission riil pada basis data. Keberhasilan penyimpanan tujuh submission internal sampai 6 Mei 2026 menjadi indikator bahwa alur dari input pengguna hingga pembacaan data oleh sistem telah berjalan secara utuh dalam satu siklus layanan.",
    )

    # Tambahkan pembahasan dan implikasi sebelum daftar pustaka.
    daftar_pustaka = find_paragraph(document, "DAFTAR PUSTAKA")
    p43 = insert_paragraph_before(daftar_pustaka, "4.3 Implikasi Tugas Akhir", "Heading 2")
    p = insert_paragraph_after(
        p43,
        "Implikasi praktis dari tugas akhir ini terletak pada tersedianya media pre-screening yang lebih transparan bagi sekolah dasar negeri dan masyarakat. Pihak sekolah memperoleh data awal calon siswa yang sudah tersusun dalam bentuk skor terukur, sedangkan orang tua atau pendamping dari TK dan PAUD memperoleh gambaran peluang kelayakan anak sebelum proses PPDB resmi berlangsung.",
        "Normal",
    )
    insert_paragraph_after(
        p,
        "Implikasi akademik penelitian ini menunjukkan bahwa integrasi metode RAD dan SAW dapat diterapkan secara efektif pada project sistem informasi pendidikan yang bersifat web-based. Hasil tugas akhir ini juga membuka ruang pengembangan lanjutan berupa penambahan data sekolah, penyempurnaan bobot kriteria melalui kebijakan sekolah, dan pelaksanaan UAT lapangan yang lebih luas untuk mematangkan kesiapan implementasi.",
        "Normal",
    )

    p42 = insert_paragraph_before(p43, "4.2 Pembahasan", "Heading 2")
    p = insert_paragraph_after(
        p42,
        "Hasil pada Bab 4 menunjukkan bahwa pemilihan metode RAD sesuai dengan karakter penelitian yang menuntut pengembangan sistem dalam waktu relatif singkat, tetapi tetap membutuhkan umpan balik bertahap dari pengguna. Tahap requirements planning menghasilkan kebutuhan yang jelas dari sekolah, tahap user design menerjemahkannya ke dalam alur antarmuka yang mudah dipahami, dan tahap construction serta cutover membuktikan bahwa kebutuhan tersebut dapat diwujudkan menjadi sistem yang berjalan pada lingkungan pengujian.",
        "Normal",
    )
    p = insert_paragraph_after(
        p,
        "Penerapan SAW juga memberikan hasil yang sejalan dengan kebutuhan lapangan karena metode ini mampu menjelaskan alasan rekomendasi secara langsung melalui bobot dan nilai normalisasi tiap kriteria. Perbedaan utama dibandingkan referensi yang menggunakan AHP terletak pada kesederhanaan proses. SAW tidak memerlukan matriks perbandingan berpasangan, sehingga lebih mudah diterapkan pada sistem sekolah dasar yang membutuhkan perhitungan cepat, transparan, dan mudah dipahami oleh admin sekolah maupun pendamping dari TK atau PAUD.",
        "Normal",
    )
    insert_paragraph_after(
        p,
        "Pembahasan pengujian memperlihatkan bahwa seluruh fitur inti telah lolos pengujian black box, sementara UAT lapangan disiapkan dengan responden yang benar-benar relevan terhadap konteks penggunaan sistem. Kondisi ini membuat hasil penelitian tetap jujur secara akademik karena tidak memaksakan angka penerimaan yang belum diukur, tetapi sekaligus menunjukkan bahwa perangkat uji sudah siap digunakan pada tahap pengambilan data lapangan berikutnya.",
        "Normal",
    )

    # Tambahkan referensi PDF baru sebelum lampiran tanpa mengganggu bibliografi Mendeley yang sudah ada.
    lampiran = find_paragraph(document, "LAMPIRAN")
    references = [
        "Beynon-Davies, P., Carne, C., Mackay, H., & Tudhope, D. (1999). Rapid application development (RAD): An empirical review. European Journal of Information Systems, 8(3), 211-223. https://damiantgordon.com/Methodologies/Papers/Rapid%20Application%20Development%20RAD%20An%20Empirical%20Review.pdf",
        "Chen, S., Thaduri, U. R., & Ballamudi, V. K. R. (2019). Front-End Development in React: An Overview. Engineering International, 7(2), 117-126. https://abc.us.org/ojs/index.php/ei/article/view/662/1191",
        "Jadhav, G., & Gonsalves, F. (2020). Role of Node.js in Modern Web Application Development. International Research Journal of Engineering and Technology, 7(6). https://www.irjet.net/archives/V7/i6/IRJET-V7I61149.pdf",
        "Kulkarni, S. T., Siddha, R., Mattani, B., Mohite, S., & Lahane, P. (2025). A Comprehensive Review of MERN Stack Development: Trends, Challenges, and Future Directions. SSRN. https://papers.ssrn.com/sol3/Delivery.cfm/SSRN_ID5837262_code3635775.pdf?abstractid=5837262&mirid=1&type=2",
        "Nurmalini, & Rahim, R. (2017). Study Approach of Simple Additive Weighting For Decision Support System. International Journal of Scientific Research in Science and Technology, 3(3), 541-544. https://ijsrst.com/paper/1096.pdf",
    ]
    anchor = lampiran
    for ref in reversed(references):
        new_paragraph = insert_paragraph_before(anchor, ref, "Normal")
        apply_reference_format(new_paragraph)
        anchor = new_paragraph

    document.save(str(OUTPUT_DOC))
    print(f"Saved revised document to: {OUTPUT_DOC}")


if __name__ == "__main__":
    main()
