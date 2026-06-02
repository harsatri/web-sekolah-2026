from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_COLOR_INDEX
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


SRC = r"C:\Users\upgra\Documents\trae_projects\WebSekolah-2026\BAB4_purposive_sampling_review.docx"
OUT = r"C:\Users\upgra\OneDrive\Documents\BAB4_notes.docx"


doc = Document(SRC)
orig = {i: doc.paragraphs[i] for i in range(len(doc.paragraphs))}

BLUE = RGBColor(0x0D, 0x47, 0xA1)
GREEN = RGBColor(0x1B, 0x5E, 0x20)
ORANGE = RGBColor(0xE6, 0x51, 0x00)
GRAY = RGBColor(0x61, 0x61, 0x61)
PURPLE = RGBColor(0x6A, 0x1B, 0x9A)


def insert_paragraph_after(paragraph, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        para.style = style
    return para


def add_line(p, label, body, color=BLUE, highlight=WD_COLOR_INDEX.TURQUOISE):
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = color
    r1.font.highlight_color = highlight
    r2 = p.add_run(body)
    r2.font.size = Pt(10)
    r2.font.color.rgb = color
    r2.font.highlight_color = highlight


def add_note(anchor, title, before, afters, color=BLUE, highlight=WD_COLOR_INDEX.TURQUOISE):
    current = anchor
    title_p = insert_paragraph_after(current, style=doc.styles["Normal"])
    add_line(
        title_p,
        f"[CATATAN BAB IV - {title}] ",
        "Usulan paragraf berikut disusun untuk melengkapi struktur penulisan tanpa mengubah bagian yang sudah benar.",
        color=color,
        highlight=highlight,
    )
    current = title_p
    before_p = insert_paragraph_after(current, style=doc.styles["Normal"])
    add_line(
        before_p,
        "[BEFORE - KONDISI SAAT INI] ",
        before,
        color=GRAY,
        highlight=WD_COLOR_INDEX.GRAY_25,
    )
    current = before_p
    for label, body in afters:
        p = insert_paragraph_after(current, style=doc.styles["Normal"])
        add_line(p, label, body, color=color, highlight=highlight)
        current = p


def diagram_note(idx, title, kind, focus):
    add_note(
        orig[idx],
        title,
        f"Subbagian ini saat ini baru berupa judul {kind.lower()} tanpa paragraf pendukung dan tanpa penjelasan sebelum maupun sesudah diagram.",
        [
            (
                "[AFTER - PARAGRAF PEMBUKA] ",
                f"Subbagian {title.lower()} menjelaskan {focus} melalui representasi {kind.lower()} pada sistem DILAYAKIN. Uraian ini diperlukan agar pembaca memahami tujuan penyusunan diagram dan kaitannya dengan proses bisnis sistem.",
            ),
            (
                "[AFTER - PARAGRAF PENGANTAR DIAGRAM] ",
                f"Diagram pada bagian ini disajikan untuk memperlihatkan alur interaksi, urutan aktivitas, atau struktur proses yang terkait dengan {title.lower()}. Penyajian visual tersebut membantu peneliti menjelaskan logika sistem secara lebih sistematis.",
            ),
            (
                "[AFTER - PARAGRAF SESUDAH DIAGRAM] ",
                f"Diagram tersebut menunjukkan bahwa proses {title.lower()} telah dirancang mengikuti kebutuhan pengguna dan aturan logika sistem yang dibangun. Hubungan antaraktor, tahapan proses, atau pertukaran data pada diagram mendukung keterpahaman pembaca terhadap mekanisme sistem.",
            ),
            (
                "[AFTER - PARAGRAF PENUTUP] ",
                f"Penjelasan pada subbagian {title.lower()} memperkuat konsistensi rancangan sistem pada tahap user design. Kejelasan diagram ini juga membantu proses implementasi pada tahap konstruksi.",
            ),
        ],
    )


add_note(
    orig[469],
    "Pengumpulan Data",
    "Subbab ini sudah memiliki paragraf pembuka, pengantar tabel, dan paragraf sesudah tabel, tetapi belum memiliki paragraf penutup yang menegaskan kontribusi hasil pengumpulan data terhadap tahap berikutnya.",
    [
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Hasil pengumpulan data tersebut memberikan dasar yang kuat bagi peneliti untuk melanjutkan proses identifikasi masalah, perumusan kebutuhan sistem, dan penyusunan parameter penilaian pada metode Simple Additive Weighting (SAW). Kesesuaian data dengan kondisi lapangan juga memperkuat relevansi sistem yang dibangun terhadap kebutuhan sekolah dasar negeri di wilayah penelitian.",
        )
    ],
    color=GREEN,
    highlight=WD_COLOR_INDEX.BRIGHT_GREEN,
)

add_note(
    orig[470],
    "Identifikasi Masalah",
    "Subbab ini saat ini hanya berupa judul tanpa paragraf pembuka, isi utama, dan penutup.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Subbab identifikasi masalah memaparkan hasil analisis peneliti terhadap kondisi penerimaan calon siswa sekolah dasar negeri yang menjadi dasar pengembangan sistem. Penjabaran masalah diperlukan agar arah perancangan sistem dapat disusun secara tepat sesuai kebutuhan pengguna dan kondisi operasional sekolah.",
        ),
        (
            "[AFTER - PARAGRAF INTI] ",
            "Hasil identifikasi masalah menunjukkan bahwa proses penilaian kelayakan calon siswa pada beberapa sekolah dasar negeri di wilayah Purwokerto masih menghadapi kendala pada aspek efisiensi, objektivitas, dan transparansi. Proses penilaian yang dilakukan secara manual menyebabkan pengolahan data membutuhkan waktu lebih lama, menghasilkan potensi perbedaan penilaian antar panitia, serta belum menyediakan dasar perhitungan yang terukur bagi orang tua maupun pihak sekolah. Kondisi tersebut menunjukkan perlunya sistem berbasis website yang mampu mengintegrasikan kriteria usia, domisili, nilai rapor TK, prestasi, kelengkapan dokumen, dan kondisi ekonomi ke dalam proses penilaian yang lebih sistematis.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Rumusan masalah yang diperoleh dari tahap ini menjadi landasan dalam penyusunan kebutuhan fungsional dan nonfungsional sistem. Arah pengembangan sistem DILAYAKIN selanjutnya difokuskan pada penyediaan sarana pre-screening yang lebih cepat, objektif, dan mudah dipahami oleh pengguna.",
        ),
    ],
)

add_note(
    orig[479],
    "Analisis Sistem Existing",
    "Subbagian ini sudah memiliki paragraf pembuka dan paragraf penjelas sesudah gambar, tetapi belum memiliki paragraf pengantar sebelum gambar, keterangan caption gambar yang jelas, dan paragraf penutup subbagian.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR GAMBAR] ",
            "Visualisasi alur sistem existing diperlukan untuk memperjelas tahapan pendaftaran, verifikasi, penilaian, dan wawancara yang selama ini dilakukan secara manual oleh pihak sekolah. Penyajian alur tersebut membantu peneliti menunjukkan titik-titik proses yang masih memerlukan perbaikan melalui sistem yang dibangun.",
        ),
        (
            "[AFTER - KETERANGAN GAMBAR] ",
            "Gambar 4.xx Alur Sistem Existing Penilaian Kelayakan Calon Siswa Sekolah Dasar Negeri",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH GAMBAR] ",
            "Gambar tersebut memperlihatkan bahwa sebagian besar aktivitas pada sistem existing masih bergantung pada input manual dan pertimbangan subjektif panitia sekolah. Pola kerja tersebut menimbulkan beban administrasi yang tinggi dan belum menghasilkan informasi kelayakan yang tersusun secara kuantitatif.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Uraian mengenai sistem existing menegaskan adanya kebutuhan terhadap mekanisme penilaian yang lebih terstandar dan berbasis perhitungan. Temuan tersebut menjadi alasan utama pengembangan sistem pre-screening berbasis website menggunakan metode SAW dalam penelitian ini.",
        ),
    ],
)

add_note(
    orig[484],
    "Analisis Sistem yang Diusulkan",
    "Subbagian ini sudah memiliki paragraf pembuka dan penjelas sesudah gambar, tetapi belum memiliki paragraf pengantar sebelum gambar, caption gambar yang jelas, dan paragraf penutup subbagian.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR GAMBAR] ",
            "Representasi alur sistem usulan disajikan untuk menunjukkan perubahan proses dari pola manual menuju proses digital yang lebih terstruktur. Alur ini juga memperlihatkan bagaimana setiap aktor berinteraksi dengan sistem dalam menghasilkan rekomendasi kelayakan calon siswa.",
        ),
        (
            "[AFTER - KETERANGAN GAMBAR] ",
            "Gambar 4.xx Alur Sistem Usulan Pre-Screening Kelayakan Calon Siswa",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH GAMBAR] ",
            "Diagram sistem usulan menunjukkan bahwa proses input data, validasi, pengolahan kriteria, dan penyajian hasil rekomendasi telah dipusatkan di dalam sistem DILAYAKIN. Struktur tersebut memungkinkan penilaian dilakukan secara lebih cepat, terukur, dan konsisten dibandingkan dengan proses existing.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Rancangan sistem yang diusulkan memperlihatkan bahwa penelitian ini tidak hanya memindahkan proses manual ke media digital, tetapi juga memperbaiki mekanisme penilaian melalui integrasi metode SAW. Arah rancangan ini menjadi dasar bagi pembahasan tahapan perhitungan pada subbagian berikutnya.",
        ),
    ],
)

add_note(
    orig[488],
    "Sub-Kriteria dan Bobot",
    "Subbagian ini saat ini langsung masuk ke rincian subkriteria tanpa paragraf pembuka yang menjelaskan fungsi keseluruhan bobot dan belum memiliki paragraf penutup setelah seluruh subkriteria selesai dipaparkan.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Subbagian subkriteria dan bobot menjelaskan dasar penyusunan parameter penilaian yang digunakan pada sistem DILAYAKIN. Penjelasan ini diperlukan agar pembaca memahami hubungan antara kebutuhan sekolah, karakteristik calon siswa, dan pembobotan yang diterapkan pada metode SAW.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Keseluruhan subkriteria dan bobot tersebut membentuk struktur penilaian yang digunakan sistem untuk menghasilkan skor kelayakan calon siswa. Kesesuaian bobot dengan kebutuhan sekolah mitra menjadi faktor penting dalam menjaga relevansi hasil rekomendasi yang dihasilkan sistem.",
        ),
    ],
)

add_note(
    orig[517],
    "Data Alternatif",
    "Subbagian ini saat ini hanya berupa judul tanpa paragraf penjelas.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Subbagian data alternatif menjelaskan bentuk data calon siswa yang digunakan sebagai masukan dalam proses perhitungan metode SAW. Data alternatif berfungsi sebagai representasi setiap calon siswa yang akan dinilai tingkat kelayakannya berdasarkan kriteria yang telah ditetapkan.",
        ),
        (
            "[AFTER - PARAGRAF INTI] ",
            "Data alternatif pada penelitian ini terdiri atas informasi identitas calon siswa beserta nilai pada setiap kriteria, meliputi usia, domisili, nilai rapor TK, prestasi, kelengkapan dokumen, dan kondisi ekonomi. Setiap alternatif direpresentasikan dalam bentuk data terstruktur agar dapat diproses pada tahap normalisasi dan perankingan. Penyusunan data alternatif secara konsisten diperlukan agar sistem mampu membandingkan tingkat kelayakan antar calon siswa secara objektif.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Penyajian data alternatif menjadi tahapan penghubung antara penentuan kriteria dengan proses perhitungan SAW. Struktur data yang jelas juga mendukung keakuratan sistem dalam menghasilkan rekomendasi sekolah dan status kelayakan calon siswa.",
        ),
    ],
)

add_note(
    orig[520],
    "Normalisasi Matriks",
    "Subbagian ini saat ini hanya menampilkan caption tabel rumus normalisasi tanpa paragraf pembuka, pengantar tabel, penjelas sesudah tabel, dan penutup.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Tahap normalisasi matriks dilakukan untuk menyetarakan nilai setiap kriteria agar dapat dibandingkan dalam satu skala perhitungan. Langkah ini diperlukan karena setiap kriteria memiliki karakteristik nilai yang berbeda, baik dalam bentuk benefit maupun cost.",
        ),
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL] ",
            "Rumus normalisasi yang digunakan pada penelitian ini disajikan untuk menunjukkan mekanisme pengubahan nilai awal setiap alternatif menjadi nilai yang siap dihitung pada tahap perankingan. Penyajian rumus ini juga membantu pembaca memahami dasar matematis yang digunakan sistem DILAYAKIN.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL] ",
            "Rumus pada tabel menunjukkan bahwa nilai setiap calon siswa diubah menjadi nilai ternormalisasi dengan mempertimbangkan jenis atribut pada masing-masing kriteria. Proses tersebut memungkinkan seluruh data dinilai secara proporsional sebelum digabungkan dengan bobot kriteria.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Hasil normalisasi matriks menjadi dasar utama dalam menghitung skor akhir kelayakan calon siswa. Ketepatan tahap ini sangat menentukan konsistensi hasil rekomendasi yang diberikan sistem.",
        ),
    ],
)

add_note(
    orig[522],
    "Perankingan",
    "Subbagian ini saat ini hanya menampilkan caption tabel rumus skor akhir tanpa paragraf pendukung.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Tahap perankingan merupakan proses akhir dalam perhitungan metode SAW yang digunakan untuk menentukan urutan kelayakan setiap calon siswa. Tahap ini menggabungkan nilai hasil normalisasi dengan bobot masing-masing kriteria.",
        ),
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL] ",
            "Rumus skor akhir disajikan untuk memperlihatkan proses akumulasi nilai ternormalisasi yang telah dikalikan dengan bobot kriteria. Hasil perhitungan ini menjadi dasar sistem dalam menetapkan tingkat kelayakan dan rekomendasi sekolah.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL] ",
            "Rumus pada tabel menunjukkan bahwa setiap calon siswa memperoleh nilai preferensi akhir berdasarkan kontribusi seluruh kriteria penilaian. Nilai preferensi yang lebih tinggi menunjukkan tingkat kelayakan yang lebih baik menurut parameter yang digunakan pada sistem.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Tahap perankingan menghasilkan dasar keputusan yang terukur dalam proses pre-screening calon siswa. Hasil pada tahap ini selanjutnya diterjemahkan ke dalam bentuk rekomendasi yang mudah dipahami oleh pengguna sistem.",
        ),
    ],
)

add_note(
    orig[523],
    "Hasil Akhir dan Rekomendasi",
    "Subbagian ini saat ini hanya berupa judul tanpa paragraf pembuka, isi, dan penutup.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Subbagian hasil akhir dan rekomendasi menjelaskan keluaran utama yang dihasilkan sistem setelah seluruh tahapan perhitungan SAW selesai dilakukan. Keluaran sistem dirancang agar dapat dipahami oleh orang tua maupun pihak sekolah sebagai dasar identifikasi awal kelayakan calon siswa.",
        ),
        (
            "[AFTER - PARAGRAF INTI] ",
            "Hasil akhir pada sistem DILAYAKIN ditampilkan dalam bentuk nilai preferensi, kategori tingkat kelayakan, dan urutan rekomendasi sekolah berdasarkan kecocokan data calon siswa terhadap kriteria yang digunakan. Penyajian hasil dilakukan secara terstruktur agar pengguna dapat mengetahui kontribusi setiap kriteria terhadap skor akhir yang diperoleh. Format keluaran ini juga memudahkan sistem dalam menyampaikan rekomendasi secara objektif dan informatif.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Penyajian hasil akhir dan rekomendasi menjadi penegas bahwa sistem yang dibangun tidak hanya melakukan perhitungan, tetapi juga menerjemahkan hasil tersebut ke dalam informasi yang dapat dimanfaatkan pengguna. Fungsi tersebut memperkuat peran sistem sebagai alat bantu pre-screening sebelum proses pendaftaran resmi dilakukan.",
        ),
    ],
)

add_note(
    orig[532],
    "Use Case Diagram",
    "Subbagian ini sudah memiliki paragraf pembuka, tetapi belum memiliki paragraf pengantar gambar atau diagram, penjelas sesudah diagram, dan penutup subbagian.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR DIAGRAM] ",
            "Use case diagram perlu disajikan untuk memperjelas ruang lingkup interaksi antara aktor dan fitur utama dalam sistem DILAYAKIN. Diagram ini membantu peneliti menunjukkan batas fungsi sistem yang dirancang pada tahap user design.",
        ),
        ("[AFTER - KETERANGAN GAMBAR] ", "Gambar 4.xx Use Case Diagram Sistem DILAYAKIN"),
        (
            "[AFTER - PARAGRAF SESUDAH DIAGRAM] ",
            "Diagram tersebut memperlihatkan bahwa setiap aktor memiliki hak akses dan fungsi yang berbeda sesuai kebutuhan operasional masing-masing. Pembagian use case tersebut menunjukkan bahwa sistem dirancang dengan struktur peran yang jelas untuk menjaga keteraturan proses bisnis.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Representasi use case diagram memperkuat hasil analisis kebutuhan pengguna yang telah diidentifikasi sebelumnya. Struktur fungsi yang tergambar pada diagram ini juga menjadi acuan dalam penyusunan diagram aktivitas, sequence diagram, dan class diagram pada bagian berikutnya.",
        ),
    ],
)

add_note(
    orig[535],
    "Activity Diagram",
    "Subbagian ini sudah memiliki paragraf pembuka, tetapi belum memiliki pengantar diagram utama dan paragraf penutup yang menjembatani ke rincian diagram aktivitas per proses.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR DIAGRAM] ",
            "Activity diagram utama perlu disajikan untuk menggambarkan aliran aktivitas sistem secara umum sebelum masuk ke rincian proses pada setiap modul. Penyajian ini membantu pembaca memahami hubungan antarproses dalam sistem secara menyeluruh.",
        ),
        ("[AFTER - KETERANGAN GAMBAR] ", "Gambar 4.xx Activity Diagram Umum Sistem DILAYAKIN"),
        (
            "[AFTER - PARAGRAF SESUDAH DIAGRAM] ",
            "Diagram aktivitas secara umum menunjukkan bahwa proses pada sistem tersusun mulai dari autentikasi pengguna, pengisian simulasi, pengolahan nilai, hingga pengelolaan data oleh admin dan super admin. Struktur tersebut menegaskan bahwa alur sistem telah dirancang secara bertahap dan terarah.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Uraian diagram aktivitas umum menjadi dasar untuk membahas diagram aktivitas yang lebih spesifik pada setiap proses utama sistem. Pembahasan rinci tersebut disajikan pada subbagian berikutnya.",
        ),
    ],
)

for idx, title, focus in [
    (536, "Login dan Autentikasi", "alur aktivitas pengguna saat mengakses sistem dan melakukan validasi akun"),
    (538, "Simulasi Kelayakan SAW", "rangkaian aktivitas simulasi kelayakan calon siswa menggunakan metode SAW"),
    (539, "Activity akses simulasi", "alur aktivitas ketika pengguna mengakses halaman simulasi kelayakan"),
    (541, "Activity input simulasi", "alur aktivitas ketika pengguna mengisi data simulasi kelayakan secara bertahap"),
    (543, "Activity perhitungan SAW dan Hasil", "alur aktivitas ketika sistem memproses data input menjadi hasil penilaian dan rekomendasi"),
    (545, "Pengajuan Kriteria oleh Admin Sekolah", "alur aktivitas admin sekolah saat mengajukan perubahan kriteria penilaian"),
    (547, "Validasi Pengajuan Kriteria oleh Super Admin", "alur aktivitas super admin saat memvalidasi pengajuan perubahan kriteria"),
]:
    diagram_note(idx, title, "diagram aktivitas", focus)

add_note(
    orig[549],
    "Sequence Diagram",
    "Subbagian ini sudah memiliki paragraf pembuka, tetapi belum memiliki pengantar diagram utama dan paragraf penutup yang menjembatani rincian urutan interaksi pada tiap proses.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR DIAGRAM] ",
            "Sequence diagram utama perlu disajikan untuk menunjukkan urutan interaksi antaraktor dan komponen sistem sebelum pembahasan diarahkan pada proses yang lebih spesifik. Diagram ini berfungsi menjelaskan pertukaran data secara kronologis dalam sistem DILAYAKIN.",
        ),
        ("[AFTER - KETERANGAN GAMBAR] ", "Gambar 4.xx Sequence Diagram Umum Sistem DILAYAKIN"),
        (
            "[AFTER - PARAGRAF SESUDAH DIAGRAM] ",
            "Diagram urutan secara umum memperlihatkan bagaimana permintaan dari pengguna diteruskan oleh antarmuka ke controller, service, dan basis data hingga menghasilkan respons yang sesuai. Pola interaksi ini memperlihatkan bahwa arsitektur sistem telah dirancang untuk mendukung pengolahan data secara terstruktur.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Penjelasan sequence diagram umum memberikan dasar konseptual sebelum pembahasan diarahkan pada urutan interaksi setiap modul. Rincian sequence diagram per proses disajikan pada subbagian berikutnya.",
        ),
    ],
)

for idx, title, focus in [
    (550, "Login dan Autentikasi", "urutan interaksi saat pengguna melakukan login dan validasi hak akses"),
    (552, "Simulasi Kelayakan SAW", "urutan interaksi utama pada proses simulasi kelayakan"),
    (553, "Sequence akses simulasi", "urutan interaksi saat pengguna membuka modul simulasi"),
    (555, "Sequence input simulasi", "urutan interaksi saat data simulasi dikirimkan ke sistem"),
    (557, "Sequence perhitungan SAW dan Hasil", "urutan interaksi saat sistem menghitung skor dan menampilkan hasil"),
    (559, "Pengajuan Kriteria oleh Admin", "urutan interaksi saat admin mengajukan perubahan kriteria"),
    (560, "Buka Halaman Pengajuan Kriteria", "urutan interaksi saat admin membuka halaman pengajuan kriteria"),
    (562, "Isi dan Kirim Form Pengajuan", "urutan interaksi saat admin mengisi dan mengirimkan form pengajuan kriteria"),
    (564, "Validasi Pengajuan Kriteria oleh Super Admin", "urutan interaksi saat super admin memeriksa dan memutuskan status pengajuan"),
]:
    diagram_note(idx, title, "sequence diagram", focus)

add_note(
    orig[567],
    "Class Diagram",
    "Subbagian ini sudah memiliki paragraf pembuka, tetapi belum memiliki pengantar diagram, penjelas sesudah diagram, dan penutup subbagian.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR DIAGRAM] ",
            "Class diagram perlu ditampilkan untuk memperjelas struktur data dan hubungan antarkelas yang mendukung proses bisnis pada sistem DILAYAKIN. Diagram ini juga membantu peneliti menunjukkan keterkaitan antara entitas pengguna, sekolah, simulasi, kriteria, dan hasil penilaian.",
        ),
        ("[AFTER - KETERANGAN GAMBAR] ", "Gambar 4.xx Class Diagram Sistem DILAYAKIN"),
        (
            "[AFTER - PARAGRAF SESUDAH DIAGRAM] ",
            "Diagram tersebut menunjukkan bahwa setiap class memiliki peran yang saling mendukung dalam mengelola data dan menjalankan logika sistem. Relasi antarkelas yang tersusun pada diagram menjadi dasar implementasi struktur basis data dan kode program pada tahap konstruksi.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Kejelasan struktur class memperkuat konsistensi rancangan sistem dari sisi data maupun logika aplikasi. Representasi ini juga memudahkan proses implementasi fitur secara modular pada tahap pengembangan.",
        ),
    ],
)

ui_context = {
    571: "autentikasi pengguna berdasarkan peran yang dimiliki dalam sistem",
    574: "akses awal pengguna terhadap informasi utama dan fitur sistem",
    577: "pengisian data simulasi kelayakan secara bertahap",
    580: "penyajian hasil perhitungan SAW dan rincian simulasi kepada pengguna",
    583: "monitoring data pengguna dan pengelolaan informasi sekolah oleh admin",
    586: "proses pengajuan perubahan kriteria oleh admin sekolah",
    588: "monitoring sistem dan pengelolaan data pada tingkat keseluruhan oleh super admin",
    591: "proses validasi pengajuan perubahan kriteria oleh super admin",
}
for idx in [571, 574, 577, 580, 583, 586, 588, 591]:
    title = orig[idx].text.strip()
    anchor = orig[idx + 1] if idx not in [586, 591] else orig[idx]
    before = (
        "Subbagian ini belum memiliki penjelasan yang lengkap mengenai fungsi halaman, pengantar gambar antarmuka, penjelas sesudah gambar, dan paragraf penutup."
        if idx in [586, 591]
        else "Subbagian ini sudah memiliki deskripsi singkat, tetapi belum memiliki pengantar gambar antarmuka, penjelas sesudah gambar, dan paragraf penutup yang menegaskan fungsi halaman."
    )
    add_note(
        anchor,
        title,
        before,
        [
            (
                "[AFTER - PARAGRAF PEMBUKA] ",
                f"Subbagian ini menjelaskan rancangan halaman {title.lower()} pada sistem DILAYAKIN yang dirancang untuk mendukung {ui_context[idx]}. Penjelasan ini diperlukan agar pembaca memahami tujuan antarmuka sebelum melihat visualisasi rancangan halaman.",
            ),
            (
                "[AFTER - PARAGRAF PENGANTAR GAMBAR] ",
                f"Gambar rancangan halaman {title.lower()} disajikan untuk menunjukkan komponen antarmuka, struktur navigasi, dan informasi utama yang tersedia bagi pengguna. Penyajian visual ini membantu peneliti menjelaskan kesesuaian rancangan dengan kebutuhan pengguna.",
            ),
            (
                "[AFTER - PARAGRAF SESUDAH GAMBAR] ",
                f"Tampilan pada halaman {title.lower()} menunjukkan bahwa sistem dirancang dengan susunan elemen yang mendukung kemudahan penggunaan dan keterbacaan informasi. Komponen yang ditampilkan juga selaras dengan peran pengguna dan fungsi modul yang diakses.",
            ),
            (
                "[AFTER - PARAGRAF PENUTUP] ",
                f"Rancangan halaman {title.lower()} memperlihatkan bahwa aspek antarmuka telah disusun untuk menunjang efektivitas penggunaan sistem. Kejelasan antarmuka ini menjadi bagian penting dalam mendukung penerimaan pengguna terhadap sistem yang dibangun.",
            ),
        ],
    )

page_use = {
    596: "proses autentikasi dan pembatasan akses pengguna",
    598: "penyajian informasi umum serta akses menuju fitur utama sistem",
    600: "proses input data simulasi kelayakan calon siswa",
    602: "penyajian nilai preferensi, status kelayakan, dan rekomendasi sekolah",
    604: "monitoring data pengguna dan hasil simulasi oleh admin sekolah",
    606: "pengajuan perubahan kriteria penilaian oleh admin sekolah",
    608: "monitoring keseluruhan sistem dan pengelolaan data utama oleh super admin",
    610: "validasi pengajuan perubahan kriteria oleh super admin",
}
for idx in [596, 598, 600, 602, 604, 606, 608, 610]:
    title = orig[idx].text.strip()
    add_note(
        orig[idx],
        title,
        "Subbagian ini saat ini hanya berupa judul tanpa paragraf pembuka, penjelas implementasi, pengantar gambar tangkapan layar, penjelas sesudah gambar, dan penutup.",
        [
            (
                "[AFTER - PARAGRAF PEMBUKA] ",
                f"Subbagian ini memaparkan hasil implementasi halaman {title.lower()} pada sistem DILAYAKIN sebagai bagian dari tahap konstruksi. Pembahasan difokuskan pada kesesuaian hasil implementasi dengan rancangan antarmuka dan kebutuhan pengguna.",
            ),
            (
                "[AFTER - PARAGRAF PENGANTAR GAMBAR] ",
                f"Tangkapan layar halaman {title.lower()} disajikan untuk menunjukkan hasil implementasi antarmuka yang telah dibangun pada sistem. Visualisasi ini membantu pembaca membandingkan hasil implementasi dengan rancangan yang telah dijelaskan pada tahap user design.",
            ),
            (
                "[AFTER - PARAGRAF SESUDAH GAMBAR] ",
                f"Tampilan hasil implementasi halaman {title.lower()} menunjukkan bahwa sistem telah mampu mendukung {page_use[idx]}. Susunan komponen pada halaman ini juga memperlihatkan keterpaduan antara fungsi, data, dan pengalaman pengguna.",
            ),
            (
                "[AFTER - PARAGRAF PENUTUP] ",
                f"Implementasi halaman {title.lower()} menunjukkan bahwa proses konstruksi telah berhasil menerjemahkan rancangan menjadi fitur yang dapat digunakan. Kesesuaian implementasi ini mendukung pencapaian tujuan sistem secara keseluruhan.",
            ),
        ],
        color=PURPLE,
        highlight=WD_COLOR_INDEX.PINK,
    )

add_note(
    orig[613],
    "Black box Testing",
    "Subbab ini saat ini langsung diikuti daftar fitur yang diuji tanpa paragraf pembuka yang menjelaskan tujuan, ruang lingkup, dan cara membaca hasil pengujian.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Pengujian black box dilakukan untuk menilai kesesuaian fungsi sistem DILAYAKIN terhadap kebutuhan yang telah dirancang pada tahap sebelumnya. Pengujian ini difokuskan pada masukan, proses, dan keluaran setiap fitur tanpa meninjau struktur kode program secara internal.",
        ),
        (
            "[AFTER - PARAGRAF PENJELAS] ",
            "Setiap tabel pengujian pada subbagian ini menyajikan skenario proses, aktor yang terlibat, hasil yang diharapkan, dan hasil aktual sistem. Penyajian tersebut digunakan untuk menunjukkan apakah masing-masing fitur telah berjalan sesuai dengan rancangan dan kebutuhan pengguna.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP SUBBAB] ",
            "Hasil pengujian black box pada seluruh fitur menjadi dasar untuk menilai kesiapan fungsional sistem sebelum dilakukan pengujian penerimaan pengguna. Temuan pada bagian ini juga memperlihatkan tingkat kestabilan implementasi fitur yang telah dibangun.",
        ),
    ],
    color=PURPLE,
    highlight=WD_COLOR_INDEX.PINK,
)

for idx in [614, 615, 616, 617, 622, 623, 624, 625, 626, 627, 628, 629, 630, 631, 632, 633, 634, 635, 636, 637, 638, 639]:
    title = orig[idx].text.strip()
    add_note(
        orig[idx],
        f"Black Box - {title}",
        "Poin ini saat ini langsung diikuti tabel pengujian tanpa paragraf pembuka, pengantar tabel, penjelas hasil sesudah tabel, dan penutup singkat.",
        [
            (
                "[AFTER - PARAGRAF PEMBUKA] ",
                f"Pengujian pada fitur {title.lower()} dilakukan untuk memastikan bahwa fungsi yang tersedia pada modul tersebut telah berjalan sesuai dengan kebutuhan operasional sistem DILAYAKIN. Fokus pengujian diarahkan pada ketepatan respons sistem terhadap skenario penggunaan yang relevan.",
            ),
            (
                "[AFTER - PARAGRAF PENGANTAR TABEL] ",
                f"Tabel pengujian pada fitur {title.lower()} menyajikan skenario proses, aktor pengguna, hasil yang diharapkan, serta hasil aktual yang diperoleh saat sistem dijalankan. Penyajian tabel ini memudahkan peneliti dalam menunjukkan validitas fungsional setiap skenario.",
            ),
            (
                "[AFTER - PARAGRAF SESUDAH TABEL] ",
                f"Hasil pada tabel menunjukkan bahwa fitur {title.lower()} telah diuji melalui beberapa skenario yang mewakili kondisi normal maupun kondisi validasi. Kesesuaian hasil aktual dengan hasil yang diharapkan menunjukkan bahwa logika fungsi pada modul ini telah berjalan secara tepat.",
            ),
            (
                "[AFTER - PARAGRAF PENUTUP] ",
                f"Pengujian pada fitur {title.lower()} memberikan bukti bahwa modul tersebut dapat mendukung proses bisnis sistem sesuai perannya. Temuan ini memperkuat kelayakan sistem dari sisi fungsi operasional.",
            ),
        ],
        color=PURPLE,
        highlight=WD_COLOR_INDEX.PINK,
    )

add_note(
    orig[642],
    "User Acceptance Testing (UAT)",
    "Subbab ini sudah memiliki paragraf pembuka dan penjelasan instrumen, tetapi belum memiliki paragraf penutup yang menegaskan alasan pemilihan responden dan fungsi UAT dalam evaluasi sistem.",
    [
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Pelaksanaan UAT pada penelitian ini diarahkan untuk memperoleh penilaian dari pengguna yang paling relevan terhadap penggunaan sistem, yaitu admin sekolah dan orang tua atau wali calon siswa. Hasil pengujian ini digunakan untuk menilai apakah sistem DILAYAKIN telah memenuhi kebutuhan pengguna dari sisi kemudahan, kegunaan, dan penerimaan umum.",
        )
    ],
    color=GREEN,
    highlight=WD_COLOR_INDEX.BRIGHT_GREEN,
)

add_note(
    orig[646],
    "Hasil UAT Admin Sekolah",
    "Subbagian ini sudah memiliki paragraf pembuka, tetapi masih memerlukan pengantar tabel yang lengkap, penjelas hasil sesudah tabel, penjelas sesudah rumus, dan paragraf penutup subbagian.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL] ",
            "Tabel hasil UAT Admin Sekolah disajikan untuk menunjukkan penilaian responden terhadap fitur-fitur yang berkaitan dengan pengelolaan data sekolah, pemantauan hasil simulasi, dan pengajuan perubahan kriteria. Data pada tabel merepresentasikan rata-rata skor dari setiap pernyataan yang dinilai oleh responden admin.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL] ",
            "Hasil penilaian pada tabel menunjukkan tingkat penerimaan Admin Sekolah terhadap fungsi-fungsi utama sistem yang digunakan pada proses operasional. Skor yang diperoleh pada setiap pernyataan menjadi dasar untuk menghitung persentase penerimaan sistem pada kelompok responden ini.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH RUMUS] ",
            "Rumus UAT yang digunakan pada bagian ini berfungsi untuk mengubah akumulasi skor responden menjadi persentase tingkat penerimaan sistem. Nilai persentase tersebut kemudian digunakan untuk menafsirkan kategori kelayakan sistem berdasarkan standar interpretasi yang telah ditetapkan.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Hasil UAT Admin Sekolah memberikan gambaran mengenai penerimaan sistem dari sisi pengelola dan pengguna internal. Penilaian ini menjadi komponen penting dalam mengevaluasi kesesuaian sistem terhadap kebutuhan operasional sekolah.",
        ),
    ],
    color=GREEN,
    highlight=WD_COLOR_INDEX.BRIGHT_GREEN,
)

add_note(
    orig[648],
    "Hasil UAT Orang Tua/User Biasa",
    "Subbagian ini saat ini hanya berupa judul yang langsung diikuti tabel tanpa paragraf pembuka, pengantar tabel, penjelas sesudah tabel, dan penutup.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Pengujian terhadap Orang Tua/Wali dilakukan untuk mengetahui tingkat penerimaan pengguna akhir terhadap sistem DILAYAKIN. Fokus pengujian diarahkan pada kemudahan penggunaan, kejelasan informasi, dan manfaat sistem dalam membantu proses simulasi kelayakan calon siswa.",
        ),
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL] ",
            "Tabel hasil UAT Orang Tua/Wali menyajikan skor penilaian terhadap pernyataan yang berkaitan dengan pengalaman penggunaan sistem dari sisi pengguna umum. Penyajian tabel ini menunjukkan bagaimana sistem dipersepsikan oleh pihak yang akan menggunakan fitur simulasi secara langsung.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL] ",
            "Hasil pada tabel memperlihatkan tingkat penerimaan Orang Tua/Wali terhadap tampilan, alur, dan keluaran informasi yang diberikan oleh sistem. Skor yang diperoleh menjadi dasar bagi peneliti untuk menilai tingkat keberterimaan sistem dari sudut pandang pengguna akhir.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Hasil UAT dari kelompok Orang Tua/Wali melengkapi evaluasi sistem yang sebelumnya diperoleh dari kelompok Admin Sekolah. Kombinasi kedua sudut pandang tersebut memberikan dasar evaluasi yang lebih menyeluruh terhadap sistem yang dibangun.",
        ),
    ],
    color=GREEN,
    highlight=WD_COLOR_INDEX.BRIGHT_GREEN,
)

add_note(
    orig[655],
    "Rekapitulasi Hasil UAT",
    "Subbagian ini sudah memiliki paragraf penjelas umum, tetapi masih memerlukan pengantar sebelum tabel kategori interpretasi, pengantar sebelum tabel rekapitulasi, penjelas sesudah tabel, dan paragraf penutup yang menegaskan makna hasil akhir.",
    [
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL KATEGORI] ",
            "Tabel kategori interpretasi persentase digunakan sebagai dasar untuk menafsirkan tingkat penerimaan sistem berdasarkan nilai yang diperoleh pada proses UAT. Acuan ini membantu peneliti dalam menghubungkan nilai persentase dengan kategori kelayakan yang sesuai.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL KATEGORI] ",
            "Kategori pada tabel interpretasi memberikan pedoman yang jelas dalam membaca nilai persentase hasil UAT. Penggunaan kategori ini memastikan bahwa penafsiran hasil dilakukan secara konsisten dan terukur.",
        ),
        (
            "[AFTER - PARAGRAF PENGANTAR TABEL REKAPITULASI] ",
            "Tabel rekapitulasi hasil UAT menyajikan ringkasan persentase penerimaan sistem dari masing-masing kelompok responden beserta kategori yang diperoleh. Penyajian ini memudahkan pembaca dalam memahami gambaran umum tingkat kelayakan sistem.",
        ),
        (
            "[AFTER - PARAGRAF SESUDAH TABEL REKAPITULASI] ",
            "Data pada tabel menunjukkan hasil akhir penerimaan sistem dari kelompok Admin Sekolah dan Orang Tua atau Wali secara terintegrasi. Nilai rata-rata yang diperoleh menjadi indikator umum untuk menilai apakah sistem telah layak digunakan sebagai alat bantu pre-screening calon siswa.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Rekapitulasi hasil UAT menegaskan bahwa evaluasi sistem tidak hanya didasarkan pada keberhasilan fungsi teknis, tetapi juga pada penerimaan pengguna. Hasil akhir pada bagian ini menjadi penegas tingkat kelayakan sistem DILAYAKIN sebelum dibahas lebih lanjut pada bagian pembahasan.",
        ),
    ],
    color=GREEN,
    highlight=WD_COLOR_INDEX.BRIGHT_GREEN,
)

add_note(
    orig[657],
    "Pembahasan",
    "Bagian ini saat ini masih berupa kalimat editorial dan belum berbentuk pembahasan akademik yang utuh.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Bagian pembahasan digunakan untuk menafsirkan hasil penelitian yang telah diperoleh pada tahap pengumpulan data, perancangan, implementasi, dan pengujian sistem. Penafsiran dilakukan dengan menempatkan hasil sistem DILAYAKIN dalam konteks kebutuhan sekolah dasar negeri dan tujuan penggunaan metode SAW pada penelitian ini.",
        ),
        (
            "[AFTER - PARAGRAF INTI] ",
            "Hasil penelitian menunjukkan bahwa penggunaan metode SAW pada sistem DILAYAKIN mampu membantu proses identifikasi awal kelayakan calon siswa melalui penggabungan beberapa kriteria yang relevan dengan kebutuhan sekolah. Bobot yang diberikan pada masing-masing kriteria mencerminkan prioritas kebijakan sekolah dan kebutuhan informasi tambahan yang belum tercakup dalam mekanisme SPMB resmi. Penekanan pada nilai rapor TK, prestasi, kelengkapan dokumen, dan kondisi ekonomi menunjukkan bahwa sistem yang dibangun berupaya menghadirkan penilaian yang lebih komprehensif daripada proses manual maupun mekanisme administratif formal semata.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Pembahasan ini menegaskan bahwa sistem yang dikembangkan memiliki kontribusi praktis dalam menyediakan dasar penilaian yang lebih terstruktur, objektif, dan informatif. Kesesuaian antara hasil perhitungan, kebutuhan pengguna, dan penerimaan sistem menjadi indikator bahwa penelitian telah mencapai tujuan yang dirumuskan sebelumnya.",
        ),
    ],
    color=ORANGE,
    highlight=WD_COLOR_INDEX.YELLOW,
)

add_note(
    orig[658],
    "Implikasi Tugas Akhir",
    "Bagian ini saat ini hanya berupa judul tanpa paragraf penjelas.",
    [
        (
            "[AFTER - PARAGRAF PEMBUKA] ",
            "Bagian implikasi tugas akhir menjelaskan manfaat hasil penelitian bagi pihak-pihak yang terkait dengan proses penerimaan calon siswa sekolah dasar negeri. Uraian implikasi diperlukan agar kontribusi sistem yang dibangun dapat dipahami secara akademik maupun praktis.",
        ),
        (
            "[AFTER - PARAGRAF INTI] ",
            "Implikasi dari penelitian ini terlihat pada tersedianya sistem berbasis website yang mampu membantu orang tua atau wali dalam melakukan simulasi kelayakan calon siswa sebelum proses pendaftaran resmi dilakukan. Sistem DILAYAKIN juga memberikan manfaat bagi sekolah karena menyediakan informasi awal mengenai profil calon siswa berdasarkan kriteria yang telah disepakati. Dari sisi pengembangan ilmu, penelitian ini menunjukkan bahwa metode SAW dapat diterapkan secara efektif pada konteks pre-screening pendidikan dasar dengan mengintegrasikan kebutuhan lokal sekolah ke dalam proses penilaian.",
        ),
        (
            "[AFTER - PARAGRAF PENUTUP] ",
            "Keberadaan implikasi tersebut menegaskan bahwa hasil tugas akhir tidak berhenti pada pembangunan aplikasi semata, tetapi juga memberikan nilai praktis bagi pengguna dan nilai akademik bagi pengembangan penelitian sejenis. Posisi tersebut memperkuat kontribusi penelitian terhadap pemanfaatan sistem pendukung keputusan pada bidang pendidikan dasar.",
        ),
    ],
    color=ORANGE,
    highlight=WD_COLOR_INDEX.YELLOW,
)

doc.save(OUT)
print(OUT)
